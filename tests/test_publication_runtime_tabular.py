from tests.test_trusted_publication_semantics import engine
"""Publication contract tests: tabular. Shared fixtures have one owner."""
from copy import deepcopy
from data_agent.agent.publication.contracts import encode
from tests.test_trusted_publication_semantics import execute
import hashlib
import json
import pytest
from tests.support.publication_runtime import _NativeProductModel, draft, request

@pytest.mark.parametrize("case",["mean_regular","mean_new","mean_gone","mean_missing","sum_zero","sum_many","sum_daily"])
def test_normal_app_contribution_observed_arithmetic_unknown_mean_scope_sse_save_refresh_export(engine,case):
    import io,hashlib
    import pandas as pd
    from tests.test_trusted_publication_native_families import _contribution_frame,_contribution_args
    from data_agent.web.app import create_app
    from data_agent.web.agent_manager import AgentManager
    from data_agent.session.history import load_session
    from data_agent.agent.publication.runtime import read_answer
    from data_agent.tools.report import export_assistant_reply
    _,cfg=engine;variant=case.split("_",1)[1];frame=_contribution_frame(variant)
    if variant=="zero":frame["amount"]=[-10.,10.,-20.,20.]
    elif variant=="many":frame=pd.DataFrame({"day":["2026-09-01"]*11+["2026-10-01"]*11,"group":["G"+str(i) for i in range(11)]*2,"amount":[float(i+1) for i in range(11)]+[float(2*(i+1)) for i in range(11)]})
    aggregation=case.split("_",1)[0];args=_contribution_args(aggregation,"experiment")
    if variant=="daily":args.update(period_a="2026-09-01~2026-09-02",period_b="2026-10-01~2026-10-04")
    params={"decomposition_dimension":"group","decomposition_aggregation":aggregation,"decomposition_date_col":"day","period_a":args["period_a"],"period_b":args["period_b"]}
    specs=[{"kind":"table","scope":{"datasets":["experiment"],"metrics":["amount"],"parameters":dict(params,decomposition_output=output)}} for output in ["decomposition","observation","period_a","period_b"]]
    specs += [{"kind":"measurement","scope":{"datasets":["experiment"],"metrics":["amount"],"quantity":"difference","parameters":dict(params,decomposition_output="total_change")}}]
    if variant=="daily":specs.append({"kind":"table","scope":{"datasets":["experiment"],"metrics":["amount"],"parameters":dict(params,decomposition_output="daily_normalized")}})
    if variant=="zero":specs.append({"kind":"measurement","scope":{"datasets":["experiment"],"metrics":["amount"],"quantity":"relative_change","parameters":dict(params,decomposition_output="total_change_pct")}})
    content="分析上传数据，将其命名为experiment；按group分解amount的两期变化，保留完整分组、实际条数、未知均值与排除范围，不作因果归因。"
    model=_NativeProductModel(content,"experiment","contribute_decomposition",args,specs)
    manager=AgentManager(client_factory=lambda **kw:model,auxiliary_client_factory=lambda:model)
    app=create_app();app.config["agent_manager"]=manager;sid="s3_cb_"+case;data=frame.to_csv(index=False).encode("utf-8")
    with app.test_client() as http:
        ticket=http.post("/api/upload",data={"file":(io.BytesIO(data),"experiment.csv")},content_type="multipart/form-data").get_json()
        assert ticket["sha256"]=="sha256:"+hashlib.sha256(data).hexdigest();model.ticket=ticket
        response=http.post("/api/chat",json={"session_id":sid,"message":content,"uploads":[ticket]})
        events=[(b.split("\n",1)[0][7:],json.loads(b.split("\ndata: ",1)[1])) for b in response.get_data(as_text=True).split("\n\n") if b.startswith("event:")]
        ends=[d for e,d in events if e=="turn_end"];assert len(ends)==1 and ends[0]["status"]=="completed",events
        refreshed=http.get("/api/sessions/"+sid).get_json()
    saved=load_session(sid)["messages"][-1];answer=read_answer(sid,saved["answer_ref"])
    assert answer["body"]==saved["content"]==ends[0]["reply_content"]
    assert answer["requirement_outcomes"][0]["status"]==("partial" if variant in {"missing","zero"} else "answered"),answer
    req=answer["source_snapshot"]["requirements"][0];claims=[c for c in answer["source_snapshot"]["claims"] if c["id"] in req["claim_ids"]]
    values={c["basis_refs"][0]["result_path"][0]:c["value"] for c in claims};rows={r["value"]:r for r in values["decomposition"]}
    expected={"mean_regular":15.,"mean_new":50.,"mean_gone":-50.,"mean_missing":15.,"sum_zero":0.,"sum_many":66.,"sum_daily":30.}[case]
    assert values["total_change"]==expected and sum(r["contribution"] for r in rows.values())==expected
    if variant=="new":assert rows["B"]["period_a_mean"] is None and rows["B"]["count_a"]==0 and rows["B"]["contribution"]==45.
    if variant=="gone":assert rows["B"]["period_b_mean"] is None and rows["B"]["count_b"]==0 and rows["B"]["contribution"]==-20.
    if variant=="many":assert set(rows)=={"G"+str(i) for i in range(11)} and all(rows["G"+str(i)]["contribution"]==i+1 for i in range(11))
    if variant=="zero":assert values["total_change_pct"] is None and all(r["contribution_pct"] is None for r in rows.values()) and "未知" in answer["body"]
    if variant=="daily":assert values["daily_normalized"]["daily_change"]==-2.5 and values["daily_normalized"]["daily_change_pct"]==-12.5
    audit=values["observation"];assert audit["source_rows"]==len(frame) and audit["excluded_rows"]==(3 if variant=="missing" else 0)
    assert values["period_a"]["rows"]==sum(r["count_a"] for r in rows.values()) and values["period_b"]["rows"]==sum(r["count_b"] for r in rows.values())
    assert "不识别因果" in answer["body"] and "abs(总变化)" in answer["body"]
    assert any(m.get("answer_ref")==saved["answer_ref"] for m in refreshed["messages"])
    exported=export_assistant_reply(sid,answer["body"],"md")
    assert answer["body"] in (cfg.sessions_resolved/sid/"reports"/exported["artifact_path"].split("/")[-1]).read_text(encoding="utf-8")
    receipt=next(r for r in manager.get(sid).context.analysis_state.tool_receipts if r["tool_name"]=="contribute_decomposition")
    assert receipt["plan_id"] and receipt["step_id"]=="native_product" and len(receipt["required_delivery_ids"])==len(specs)


@pytest.mark.parametrize("case",["regular","constant","collinear","short","missing","excluded","scope"])
def test_normal_app_factor_original_ols_unknown_and_scope_sse_save_refresh_export(engine,case):
    import io,hashlib,math
    from tests.test_trusted_publication_native_families import _factor_frame,_factor_args
    from data_agent.web.app import create_app
    from data_agent.web.agent_manager import AgentManager
    from data_agent.session.history import load_session
    from data_agent.agent.publication.runtime import read_answer
    from data_agent.tools.report import export_assistant_reply
    _,cfg=engine;frame=_factor_frame(case);args=_factor_args("experiment",unit_col="unit",time_col="time")
    features=["x","z","constant","text"] if case=="excluded" else ["x","z"];args["features"]=",".join(features)
    params={"factor_features":features,"factor_unit_col":"unit","factor_time_col":"time"}
    outputs=["estimates","observation","unit_observation","time_observation"]
    if case=="collinear":outputs.append("bivariate_associations")
    if case=="excluded":outputs.append("excluded_features")
    specs=[{"kind":"table","scope":{"datasets":["experiment"],"metrics":["amount"],"parameters":dict(params,factor_output=output)}} for output in outputs]
    content="分析上传数据，将其命名为experiment；用原按行OLS和HC3交付amount与x、z的调整关联、真实筛样及单位时间范围，未知推断保留，不作因果归因。"
    model=_NativeProductModel(content,"experiment","attribution_analysis",args,specs)
    manager=AgentManager(client_factory=lambda **kw:model,auxiliary_client_factory=lambda:model)
    app=create_app();app.config["agent_manager"]=manager;sid="s3_factor_"+case;data=frame.to_csv(index=False).encode("utf-8")
    with app.test_client() as http:
        ticket=http.post("/api/upload",data={"file":(io.BytesIO(data),"experiment.csv")},content_type="multipart/form-data").get_json()
        assert ticket["sha256"]=="sha256:"+hashlib.sha256(data).hexdigest();model.ticket=ticket
        response=http.post("/api/chat",json={"session_id":sid,"message":content,"uploads":[ticket]})
        events=[(b.split("\n",1)[0][7:],json.loads(b.split("\ndata: ",1)[1])) for b in response.get_data(as_text=True).split("\n\n") if b.startswith("event:")]
        ends=[d for e,d in events if e=="turn_end"];assert len(ends)==1 and ends[0]["status"]=="completed",events
        refreshed=http.get("/api/sessions/"+sid).get_json()
    saved=load_session(sid)["messages"][-1];answer=read_answer(sid,saved["answer_ref"])
    assert answer["body"]==saved["content"]==ends[0]["reply_content"]
    assert answer["requirement_outcomes"][0]["status"]==("partial" if case in {"constant","collinear","short","missing","excluded"} else "answered"),answer
    req=answer["source_snapshot"]["requirements"][0];claims=[c for c in answer["source_snapshot"]["claims"] if c["id"] in req["claim_ids"]]
    values={c["basis_refs"][0]["result_path"][0]:c["value"] for c in claims}
    if case=="regular":assert {r["feature"]:r["coefficient"] for r in values["estimates"]}=={"x":round(2/math.sqrt(14),8),"z":round(3/math.sqrt(14),8)}
    if case=="constant":assert all(r["coefficient"] is None and r["p_value_holm"] is None and r["reliable"] is None for r in values["estimates"]) and "未知" in answer["body"]
    if case=="collinear":assert len(values["bivariate_associations"])==2
    if case=="excluded":assert set(values["excluded_features"])=={"constant","text"}
    audit=values["observation"];assert audit["source_rows"]==len(frame) and audit["excluded_rows"]==(3 if case=="missing" else 0)
    assert values["unit_observation"]["used_for_grouping"] is False and values["time_observation"]["used_for_correlation_correction"] is False
    if case=="scope":assert values["time_observation"]["invalid_rows"]==1 and values["time_observation"]["missing_rows"]==1 and values["unit_observation"]["missing_rows"]==1
    assert "不识别因果" in answer["body"] and "未按单位聚类" in answer["body"] and "独立性仍未验证" in answer["body"]
    assert "NaN" not in answer["body"] and "Infinity" not in answer["body"]
    assert any(m.get("answer_ref")==saved["answer_ref"] for m in refreshed["messages"])
    exported=export_assistant_reply(sid,answer["body"],"md")
    assert answer["body"] in (cfg.sessions_resolved/sid/"reports"/exported["artifact_path"].split("/")[-1]).read_text(encoding="utf-8")
    receipt=next(r for r in manager.get(sid).context.analysis_state.tool_receipts if r["tool_name"]=="attribution_analysis")
    assert receipt["plan_id"] and receipt["step_id"]=="native_product" and len(receipt["required_delivery_ids"])==len(specs)


@pytest.mark.parametrize("case",["linear","zero","constant","missing","irregular","many","short"])
def test_normal_app_time_series_original_windows_heuristics_scope_sse_save_refresh_export(engine,case):
    import io,hashlib
    from tests.test_trusted_publication_native_families import _time_series_frame,_time_series_args
    from data_agent.web.app import create_app
    from data_agent.web.agent_manager import AgentManager
    from data_agent.session.history import load_session
    from data_agent.agent.publication.runtime import read_answer
    from data_agent.tools.report import export_assistant_reply
    _,cfg=engine;frame=_time_series_frame(case);args=_time_series_args("experiment")
    params={"time_series_date_col":"day","time_series_value_col":"amount"}
    outputs=["statistics","trend","observation","date_range"]
    if len(frame)>=10:outputs.extend(["window_comparisons","change_points"])
    if len(frame)>=14:outputs.extend(["autocorrelation","seasonality"])
    specs=[{"kind":"table","scope":{"datasets":["experiment"],"metrics":["amount"],"parameters":dict(params,time_series_output=output)}} for output in outputs]
    content="分析上传数据，将其命名为experiment；按原逐记录算法交付amount的时间统计、趋势和完整滑窗、自相关候选，公开有效范围与未知值，不作季节性或预测验证。"
    model=_NativeProductModel(content,"experiment","analyze_time_series",args,specs)
    manager=AgentManager(client_factory=lambda **kw:model,auxiliary_client_factory=lambda:model)
    app=create_app();app.config["agent_manager"]=manager;sid="s3_timeseries_"+case;data=frame.to_csv(index=False).encode("utf-8")
    with app.test_client() as http:
        ticket=http.post("/api/upload",data={"file":(io.BytesIO(data),"experiment.csv")},content_type="multipart/form-data").get_json()
        assert ticket["sha256"]=="sha256:"+hashlib.sha256(data).hexdigest();model.ticket=ticket
        response=http.post("/api/chat",json={"session_id":sid,"message":content,"uploads":[ticket]})
        events=[(b.split("\n",1)[0][7:],json.loads(b.split("\ndata: ",1)[1])) for b in response.get_data(as_text=True).split("\n\n") if b.startswith("event:")]
        ends=[d for e,d in events if e=="turn_end"];assert len(ends)==1 and ends[0]["status"]=="completed",events
        refreshed=http.get("/api/sessions/"+sid).get_json()
    saved=load_session(sid)["messages"][-1];answer=read_answer(sid,saved["answer_ref"])
    assert answer["body"]==saved["content"]==ends[0]["reply_content"]
    assert answer["requirement_outcomes"][0]["status"]==("partial" if case in {"zero","constant","missing","short"} else "answered"),answer
    req=answer["source_snapshot"]["requirements"][0];claims=[c for c in answer["source_snapshot"]["claims"] if c["id"] in req["claim_ids"]]
    values={c["basis_refs"][0]["result_path"][0]:c["value"] for c in claims}
    audit=values["observation"];assert audit["source_rows"]==len(frame) and audit["excluded_rows"]==(4 if case=="missing" else 0)
    if case=="linear":assert values["statistics"]["mean"]==15.5 and values["trend"]["slope"]==1. and len(values["autocorrelation"])==15
    if case=="zero":
        mid=next(r for r in values["window_comparisons"] if r["index"]==15)
        assert mid["absolute_change"]==10. and mid["change_pct"] is None and mid["threshold_exceeded"] is None
    if case=="many":assert len(values["window_comparisons"])==len(values["change_points"])==24 and {r["index"] for r in values["change_points"]}==set(range(8,32))
    if case=="constant":assert values["seasonality"]["autocorrelation"] is None and values["trend"]["p_value"] is None
    if case=="short":assert values["trend"]["slope"] is None and values["statistics"]["mean"]==1.5
    if case=="irregular":assert values["trend"]["slope"]==1. and audit["duplicate_date_rows"]==1 and audit["unique_dates"]==3
    if "seasonality" in values:assert values["seasonality"]["seasonality_validated"] is False
    assert "不证明季节性" in answer["body"] and "每条排序观测" in answer["body"] and "未校正时间相关" in answer["body"]
    assert "NaN" not in answer["body"] and "Infinity" not in answer["body"]
    assert any(m.get("answer_ref")==saved["answer_ref"] for m in refreshed["messages"])
    exported=export_assistant_reply(sid,answer["body"],"md")
    assert answer["body"] in (cfg.sessions_resolved/sid/"reports"/exported["artifact_path"].split("/")[-1]).read_text(encoding="utf-8")
    receipt=next(r for r in manager.get(sid).context.analysis_state.tool_receipts if r["tool_name"]=="analyze_time_series")
    assert receipt["plan_id"] and receipt["step_id"]=="native_product" and len(receipt["required_delivery_ids"])==len(specs)


@pytest.mark.parametrize("case",["hourly","single","missing","duplicates","short","irregular"])
def test_normal_app_forecast_original_points_unknown_bands_scope_sse_save_refresh_export(engine,case):
    import io,hashlib
    import pandas as pd
    from tests.test_trusted_publication_native_families import _forecast_frame,_forecast_args
    from data_agent.web.app import create_app
    from data_agent.web.agent_manager import AgentManager
    from data_agent.session.history import load_session
    from data_agent.agent.publication.runtime import read_answer
    from data_agent.tools.report import export_assistant_reply
    _,cfg=engine;frame=_forecast_frame(case);args=_forecast_args("experiment",case=case)
    params={"forecast_periods":args["periods"],"forecast_date_col":"day","forecast_method":"auto"}
    outputs=["forecast","observation"] if case in {"short","irregular"} else ["forecast","backtest","interval","observation","date_range"]
    specs=[{"kind":"forecast" if output=="forecast" else "table","scope":{"datasets":["experiment"],"metrics":["amount"],"parameters":dict(params,forecast_output=output)}} for output in outputs]
    content="分析上传数据，将其命名为experiment；按原auto有序留出末值/线性候选交付amount未来点预测、真实有效范围及残差启发带，保留完整时间戳和未知边界，不声称95%预测覆盖。"
    model=_NativeProductModel(content,"experiment","forecast",args,specs)
    manager=AgentManager(client_factory=lambda **kw:model,auxiliary_client_factory=lambda:model)
    app=create_app();app.config["agent_manager"]=manager;sid="s3_forecast46_"+case;data=frame.to_csv(index=False).encode("utf-8")
    with app.test_client() as http:
        ticket=http.post("/api/upload",data={"file":(io.BytesIO(data),"experiment.csv")},content_type="multipart/form-data").get_json()
        assert ticket["sha256"]=="sha256:"+hashlib.sha256(data).hexdigest();model.ticket=ticket
        response=http.post("/api/chat",json={"session_id":sid,"message":content,"uploads":[ticket]})
        events=[(b.split("\n",1)[0][7:],json.loads(b.split("\ndata: ",1)[1])) for b in response.get_data(as_text=True).split("\n\n") if b.startswith("event:")]
        ends=[d for e,d in events if e=="turn_end"];assert len(ends)==1 and ends[0]["status"]=="completed",events
        refreshed=http.get("/api/sessions/"+sid).get_json()
    saved=load_session(sid)["messages"][-1];answer=read_answer(sid,saved["answer_ref"])
    assert answer["body"]==saved["content"]==ends[0]["reply_content"]
    assert answer["requirement_outcomes"][0]["status"]==("partial" if case in {"single","missing","short","irregular"} else "answered"),answer
    req=answer["source_snapshot"]["requirements"][0];claims=[c for c in answer["source_snapshot"]["claims"] if c["id"] in req["claim_ids"]]
    values={c["basis_refs"][0]["result_path"][0]:c["value"] for c in claims}
    audit=values["observation"];assert audit["source_rows"]==len(frame) and audit["excluded_rows"]==(5 if case=="missing" else 0)
    if case=="hourly":
        assert [r["yhat"] for r in values["forecast"]]==[31.,32.,33.]
        assert [pd.Timestamp(r["date"]) for r in values["forecast"]]==list(pd.date_range("2026-09-02 06:00:00",periods=3,freq="h"))
    if case=="duplicates":assert [r["yhat"] for r in values["forecast"]]==[62.,64.,66.] and audit["collapsed_duplicate_rows"]==30
    if case=="single":
        assert values["forecast"][0]["yhat"]==round(float(frame.amount.iloc[-1]),8)
        assert values["forecast"][0]["yhat_lower"] is values["forecast"][0]["yhat_upper"] is None
    if case in {"short","irregular"}:assert "forecast" not in values
    else:
        assert values["interval"]["coverage_probability"] is None and values["interval"]["future_uncertainty_validated"] is False
        assert "未校准" in answer["body"] and "不表示95%预测覆盖" in answer["body"]
    assert "NaN" not in answer["body"] and "Infinity" not in answer["body"]
    assert any(m.get("answer_ref")==saved["answer_ref"] for m in refreshed["messages"])
    exported=export_assistant_reply(sid,answer["body"],"md")
    assert answer["body"] in (cfg.sessions_resolved/sid/"reports"/exported["artifact_path"].split("/")[-1]).read_text(encoding="utf-8")
    receipt=next(r for r in manager.get(sid).context.analysis_state.tool_receipts if r["tool_name"]=="forecast")
    assert receipt["plan_id"] and receipt["step_id"]=="native_product" and len(receipt["required_delivery_ids"])==len(specs)


@pytest.mark.parametrize("case",["regular","missing","nonfinite","unknown","zero","invalid_date","new_group","ratio","ratio_missing"])
def test_normal_app_period_comparison_real_sums_unknown_scope_sse_save_refresh_export(engine,case):
    import io,hashlib
    from tests.test_trusted_publication_native_families import _periods_frame,_periods_args
    from data_agent.web.app import create_app
    from data_agent.web.agent_manager import AgentManager
    from data_agent.session.history import load_session
    from data_agent.agent.publication.runtime import read_answer
    from data_agent.tools.report import export_assistant_reply
    _,cfg=engine;frame=_periods_frame("regular" if case.startswith("ratio") else case);args=_periods_args("experiment")
    if case=="new_group":args["dimensions"]="group"
    if case.startswith("ratio"):
        frame["orders"]=1.
        if case=="ratio_missing":frame.loc[0,"orders"]=None
        args.update(metrics="cost_per_order",metric_formulas={"cost_per_order":{"operator":"ratio_of_sums","numerator":"amount","denominator":"orders"}})
    metric=args["metrics"];params={"comparison_date_col":"day","comparison_dimensions":["group"] if case=="new_group" else [],"comparison_formulas":args.get("metric_formulas",{}),"period_a":args["period_a"],"period_b":args["period_b"]}
    specs=[{"kind":"table","scope":{"datasets":["experiment"],"metrics":[metric],"parameters":dict(params,comparison_output=output)}} for output in ["comparison","observation"]]
    content="分析上传数据，将其命名为experiment；比较2026-09-01至2026-09-02与2026-09-03至2026-09-04两个含首尾日期窗口的原数值合计、差值、日均或显式窗口比率，交付真实各期各组范围和未知值，不把未观测补零，不声称显著或因果。"
    model=_NativeProductModel(content,"experiment","compare_periods",args,specs)
    manager=AgentManager(client_factory=lambda **kw:model,auxiliary_client_factory=lambda:model)
    app=create_app();app.config["agent_manager"]=manager;sid="s3_periods48_"+case;data=frame.to_csv(index=False).encode("utf-8")
    with app.test_client() as http:
        ticket=http.post("/api/upload",data={"file":(io.BytesIO(data),"experiment.csv")},content_type="multipart/form-data").get_json()
        assert ticket["sha256"]=="sha256:"+hashlib.sha256(data).hexdigest();model.ticket=ticket
        response=http.post("/api/chat",json={"session_id":sid,"message":content,"uploads":[ticket]})
        events=[(b.split("\n",1)[0][7:],json.loads(b.split("\ndata: ",1)[1])) for b in response.get_data(as_text=True).split("\n\n") if b.startswith("event:")]
        ends=[d for e,d in events if e=="turn_end"];assert len(ends)==1 and ends[0]["status"]=="completed",events
        refreshed=http.get("/api/sessions/"+sid).get_json()
    saved=load_session(sid)["messages"][-1];answer=read_answer(sid,saved["answer_ref"])
    assert answer["body"]==saved["content"]==ends[0]["reply_content"]
    assert answer["requirement_outcomes"][0]["status"]==("answered" if case in {"regular","ratio"} else "partial"),answer
    req=answer["source_snapshot"]["requirements"][0];claims=[c for c in answer["source_snapshot"]["claims"] if c["id"] in req["claim_ids"]]
    values={c["basis_refs"][0]["result_path"][0]:c["value"] for c in claims}
    audit=values["observation"];assert audit["source_rows"]==len(frame) and audit["date_excluded_rows"]==(2 if case=="invalid_date" else 0)
    if case=="new_group":
        item=next(r for r in values["comparisons"] if r["dimension"]=="B")[metric]
        assert item["exact_aggregates"]=={"a":None,"b":40.,"diff":None} and item["observation"]["a"]["source_rows"]==0
    else:
        item=values["metrics"][metric];expected={"regular":30.,"missing":20.,"nonfinite":20.,"unknown":None,"zero":0.,"invalid_date":30.,"ratio":15.,"ratio_missing":None}[case]
        assert item["exact_aggregates"]["period_a"]==expected and item["exact_aggregates"]["period_b"]==(35. if case.startswith("ratio") else 70.)
        if not case.startswith("ratio"):assert item["daily_avg_b"]==35. and item["observation"]["period_b"]["effective_rows"]==2
        else:assert item["aggregation"]["operand_sums"]["period_b"]=={"numerator":70.,"denominator":2.}
    assert "描述性算术" in answer["body"] and "不授予因果或预测验证" in answer["body"]
    assert "NaN" not in answer["body"] and "Infinity" not in answer["body"]
    assert any(m.get("answer_ref")==saved["answer_ref"] for m in refreshed["messages"])
    exported=export_assistant_reply(sid,answer["body"],"md")
    assert answer["body"] in (cfg.sessions_resolved/sid/"reports"/exported["artifact_path"].split("/")[-1]).read_text(encoding="utf-8")
    receipt=next(r for r in manager.get(sid).context.analysis_state.tool_receipts if r["tool_name"]=="compare_periods")
    assert receipt["plan_id"] and receipt["step_id"]=="native_product" and len(receipt["required_delivery_ids"])==2


@pytest.mark.parametrize("tool,case,compact",[("quick_profile","unused_missing",False),("quick_profile","infinite",True),("assess_readiness","infinite",False),("assess_readiness","regular",False)])
def test_normal_app_diagnostic_actual_columns_method_boundary_sse_save_refresh_export(engine,tool,case,compact):
    import io,hashlib,pandas as pd
    from data_agent.web.app import create_app
    from data_agent.web.agent_manager import AgentManager
    from data_agent.session.history import load_session
    from data_agent.agent.publication.runtime import read_answer
    from data_agent.tools.report import export_assistant_reply
    _,cfg=engine;frame=pd.DataFrame({"amount":[1.,2.,3.]})
    if case=="infinite":frame.loc[1,"amount"]=float("inf")
    if case=="unused_missing":frame["unused"]=None
    args={"name":"experiment",**({"compact":compact} if tool=="quick_profile" else {"intent":"classification"})}
    params={"profile_compact":compact} if tool=="quick_profile" else {"readiness_intent":"classification"}
    status_key="readiness" if tool=="quick_profile" else "overall"
    specs=[{"kind":"table","scope":{"datasets":["experiment"],"parameters":dict(params,diagnostic_output=key)}} for key in ["column_diagnostics",status_key]]
    content="请载入上传数据为experiment，交付列级基础诊断和检查状态；无关列问题不能阻断独立分析，诊断不代表任何方法前提已验证。"
    model=_NativeProductModel(content,"experiment",tool,args,specs);manager=AgentManager(client_factory=lambda **kw:model,auxiliary_client_factory=lambda:model)
    app=create_app();app.config["agent_manager"]=manager;sid="s3_diagnostic49_"+tool+"_"+case;data=frame.to_csv(index=False).encode("utf-8")
    with app.test_client() as http:
        ticket=http.post("/api/upload",data={"file":(io.BytesIO(data),"experiment.csv")},content_type="multipart/form-data").get_json()
        assert ticket["sha256"]=="sha256:"+hashlib.sha256(data).hexdigest();model.ticket=ticket
        response=http.post("/api/chat",json={"session_id":sid,"message":content,"uploads":[ticket]})
        events=[(b.split("\n",1)[0][7:],json.loads(b.split("\ndata: ",1)[1])) for b in response.get_data(as_text=True).split("\n\n") if b.startswith("event:")]
        ends=[d for e,d in events if e=="turn_end"];assert len(ends)==1 and ends[0]["status"]=="completed",events
        refreshed=http.get("/api/sessions/"+sid).get_json()
    saved=load_session(sid)["messages"][-1];answer=read_answer(sid,saved["answer_ref"])
    assert answer["body"]==saved["content"]==ends[0]["reply_content"]
    assert answer["requirement_outcomes"][0]["status"]=="answered",answer
    req=answer["source_snapshot"]["requirements"][0];claims=[c for c in answer["source_snapshot"]["claims"] if c["id"] in req["claim_ids"]];values={c["basis_refs"][0]["result_path"][0]:c["value"] for c in claims}
    assert values[status_key]=="ready_with_warnings"
    assert values["column_diagnostics"]["amount"]["source_rows"]==3
    assert values["column_diagnostics"]["amount"]["nonfinite_count"]==(1 if case=="infinite" else 0)
    if case=="unused_missing":assert values["column_diagnostics"]["unused"]["missing_count"]==3
    assert "具体可行性由所选方法" in answer["body"] and "不授予统计、预测或因果验证" in answer["body"]
    assert any(m.get("answer_ref")==saved["answer_ref"] for m in refreshed["messages"])
    exported=export_assistant_reply(sid,answer["body"],"md")
    assert answer["body"] in (cfg.sessions_resolved/sid/"reports"/exported["artifact_path"].split("/")[-1]).read_text(encoding="utf-8")
    receipt=next(r for r in manager.get(sid).context.analysis_state.tool_receipts if r["tool_name"]==tool)
    assert receipt["plan_id"] and receipt["step_id"]=="native_product" and len(receipt["required_delivery_ids"])==2


def test_rehearsal_reads_exact_shared_business_and_delivery_scopes_from_public_view():
    from scripts.testing.run_p7_browser_rehearsal import ScriptedClient
    scopes=[{"datasets":["original"],"metrics":["amount"]},
            {"datasets":["original"],"parameters":{"period_a":"2026-09-01~2026-09-02","period_b":"2026-09-03~2026-09-04"},"quantity":"sum"}]
    view={"snapshot_digest":"frozen","requirements":[],"claims":[],"artifacts":[],
          "model_view_schema":"publication_model_view.v1","catalog_context":{"scopes":scopes},
          "execution_targets":[{"id":"original_goal","business_scope_ref":0,"necessary_deliveries":[{"id":"original_delivery","scope_ref":1}]}]}
    before=deepcopy(view)
    decoded=ScriptedClient._current_host_snapshot([],"</accepted_answer_contract>\n"+encode(view))
    target=decoded["execution_targets"][0]
    assert target["business_scope"]==scopes[0] and target["necessary_deliveries"][0]["scope"]==scopes[1]
    assert "business_scope_ref" not in target and "scope_ref" not in target["necessary_deliveries"][0]
    assert decoded["snapshot_digest"]==view["snapshot_digest"] and view==before


@pytest.mark.parametrize("fmt,replace,fault", [(f, r, "") for f in ["csv", "tsv", "excel", "json", "jsonl"] for r in [False, True]] + [("csv", True, f) for f in ["clean", "derive", "persist", "meta", "trust", "state"]])
def test_normal_app_file_formats_reload_before_plan_preserve_exact_sources_and_delivery(engine, monkeypatch, fmt, replace, fault):
    import io
    import pandas as pd
    from data_agent.web.app import create_app
    from data_agent.web.agent_manager import AgentManager
    from data_agent.session.history import load_session
    from data_agent.agent.publication.runtime import read_answer
    from data_agent.tools.report import export_assistant_reply
    from data_agent.llm.client import Response, ToolCall
    _, cfg = engine
    suffix = {"excel": "xlsx"}.get(fmt, fmt)
    def payload(frame):
        if fmt == "excel":
            stream = io.BytesIO(); frame.to_excel(stream, index=False); return stream.getvalue()
        if fmt in {"csv", "tsv"}: return frame.to_csv(index=False, sep="\t" if fmt == "tsv" else ",").encode("utf-8")
        return frame.to_json(orient="records", lines=fmt == "jsonl", force_ascii=False).encode("utf-8")
    data = payload(pd.DataFrame({"amount": [11, 29], "label": ["甲", "乙"]}))
    second = cfg.workspace_resolved / "inbox" / ("replacement." + suffix)
    second.parent.mkdir(parents=True, exist_ok=True)
    second_bytes = payload(pd.DataFrame({"amount": [7, 13, 22], "label": ["新甲", "新乙", "新丙"]}))
    second.write_bytes(second_bytes)
    content = "载入上传数据，将其命名为experiment；报告当前实际记录数。"
    class FileModel(_NativeProductModel):
        phase = 0
        def chat(self, messages, system="", **kw):
            if self.phase == 1 and system == "你是一个JSON输出器，只返回JSON，不要其他内容。":
                public = messages[-1]["content"]
                declared, _ = json.JSONDecoder().raw_decode(public.partition("canonical_requirements: ")[2])
                prior = next(r for r in reversed(declared) if r["lifecycle_status"] == "active" and r["semantic_kind"] == "result")
                return Response(text=encode({"intent_type": "directed_analysis", "requested_outputs": [{"exact_span": self.content, "kind": "result", "op": "replace", "target_id": prior["id"], "business_scope": {"datasets": ["experiment"], "quantity": "rows"}}]}))
            if fault and self.phase == 1 and self.step == 2 and system != "你是一个JSON输出器，只返回JSON，不要其他内容。":
                from scripts.testing.run_p7_browser_rehearsal import ScriptedClient
                view = ScriptedClient._current_host_snapshot(messages, system)
                assert "replacement_source_unverified" in messages[-1]["content"]
                return Response(text=encode({"schema": "answer_bindings.v2", "snapshot_digest": view["snapshot_digest"], "sections": []}))
            if self.phase == 1 and system != "你是一个JSON输出器，只返回JSON，不要其他内容。" and self.step == 0 and not getattr(self, "reloaded", False):
                self.reloaded = True
                self.calls.append("reload_data")
                return Response(tool_calls=[ToolCall("explicit_reload", "reload_data", {"name": "experiment", "source": "replacement." + suffix})])
            if system != "你是一个JSON输出器，只返回JSON，不要其他内容。" and self.step == 0:
                self.step = 1; self.calls.append("quick_profile")
                return Response(tool_calls=[ToolCall("file_schema", "quick_profile", {"name": "experiment"})])
            return super().chat(messages, system=system, **kw)
    specs = [{"kind": "measurement", "scope": {"datasets": ["experiment"], "quantity": "rows"}}]
    model = FileModel(content, "experiment", "describe_dataset", {"name": "experiment"}, specs)
    manager = AgentManager(client_factory=lambda **kw: model, auxiliary_client_factory=lambda: model)
    app = create_app(); app.config["agent_manager"] = manager
    sid = "s3_file_" + fmt + ("_reload" if replace else "_load")
    with app.test_client() as http:
        ticket = http.post("/api/upload", data={"file": (io.BytesIO(data), "experiment." + suffix)}, content_type="multipart/form-data").get_json()
        assert ticket["sha256"] == "sha256:" + hashlib.sha256(data).hexdigest(); model.ticket = ticket
        response = http.post("/api/chat", json={"session_id": sid, "message": content, "uploads": [ticket]})
        events = [(b.split("\n", 1)[0][7:], json.loads(b.split("\ndata: ", 1)[1])) for b in response.get_data(as_text=True).split("\n\n") if b.startswith("event:")]
        ends = [d for e, d in events if e == "turn_end"]
        assert len(ends) == 1 and ends[0]["status"] == "completed", events
        prior_saved = load_session(sid)["messages"][-1]
        prior_answer = read_answer(sid, prior_saved["answer_ref"])
        assert prior_answer["requirement_outcomes"][0]["status"] == "answered"
        if replace:
            model.phase = 1; model.step = 0
            model.content = "明确加载replacement新版并替换原目标，报告experiment新版实际记录数。"
            if fault:
                import data_agent.tools.data_clean as clean
                from data_agent.session.workspace import Workspace
                from data_agent.tools import data_io
                from data_agent.agent.analysis_state import AnalysisSessionState
                owner, method = ((clean, "auto_clean") if fault == "clean" else
                                 (Workspace, "derive") if fault == "derive" else
                                 (Workspace, "persist_dataset") if fault == "persist" else
                                 (Workspace, "save_meta") if fault == "meta" else
                                 (data_io, "_persist_trust_record") if fault == "trust" else
                                 (AnalysisSessionState, "save"))
                original = getattr(owner, method)
                def fail(*a, **kw):
                    if fault in {"persist", "meta", "state"}:
                        import inspect
                        if not any(f.function == "commit" and f.filename.endswith("data_io.py") for f in inspect.stack()):
                            return original(*a, **kw)
                    if fault != "clean": original(*a, **kw)
                    raise RuntimeError("injected app ingestion " + fault)
                monkeypatch.setattr(owner, method, fail)
            response = http.post("/api/chat", json={"session_id": sid, "message": model.content})
            events = [(b.split("\n", 1)[0][7:], json.loads(b.split("\ndata: ", 1)[1])) for b in response.get_data(as_text=True).split("\n\n") if b.startswith("event:")]
            ends = [d for e, d in events if e == "turn_end"]
            assert len(ends) == 1 and ends[0]["status"] == "completed", events
            assert read_answer(sid, prior_saved["answer_ref"]) == prior_answer
        refreshed = http.get("/api/sessions/" + sid).get_json()
    saved = load_session(sid)["messages"][-1]; answer = read_answer(sid, saved["answer_ref"])
    assert answer["body"] == saved["content"] == ends[0]["reply_content"]
    req = next(r for r in answer["source_snapshot"]["requirements"] if not r["withdrawn"])
    assert next(o for o in answer["requirement_outcomes"] if o["requirement_id"] == req["id"])["status"] == ("evidence_insufficient" if fault else "answered"), answer
    assert [c["value"] for c in answer["source_snapshot"]["claims"] if c["id"] in req["claim_ids"]] == ([] if fault else [3 if replace else 2])
    from data_agent.agent.context import use_agent_context
    with use_agent_context(manager.get(sid).context):
        ws = manager.get(sid).context.workspace
        identity = ws._publication_data_identity("experiment")
        assert identity["source_origins"][0]["file_sha256"] == "sha256:" + hashlib.sha256(second_bytes if replace and not fault else data).hexdigest()
        meta = json.loads((cfg.sessions_resolved / sid / "workspace_meta.json").read_text(encoding="utf-8"))
        parents = [item["data_identity"] for item in meta.values() if item["data_identity"]["version_id"] in identity["parent_version_ids"]]
        assert len(parents) == 1 and parents[0]["role"] == "raw"
        assert parents[0]["source_origins"] == identity["source_origins"]
        if replace and not fault:
            assert ws._publication_data_identity("experiment__history")["source_origins"][0]["file_sha256"] == ticket["sha256"]
            assert ws._publication_data_identity("experiment__history")["version_id"] != identity["version_id"]
    chain = ["quick_profile", "record_analysis_plan", "describe_dataset"]
    assert model.calls == chain + (["reload_data"] + (chain[:2] if fault else chain) if replace else [])
    if replace:
        active = next(r for r in manager.get(sid).context.analysis_state.answer_requirements if r["id"] == req["id"])
        assert active["source_binding"]["status"] == ("pending" if fault else "frozen")
        if fault:
            assert not active["source_binding"].get("verified_load_receipts")
            with use_agent_context(manager.get(sid).context):
                assert not ws._publication_data_identity("experiment__history")
            assert "新版来源尚未成功加载核验" in answer["body"]
        else: assert active["source_binding"]["verified_load_receipts"]["experiment"]
    receipt = next(r for r in reversed(manager.get(sid).context.analysis_state.tool_receipts) if r["tool_name"] == "describe_dataset")
    assert receipt["plan_id"] and receipt["step_id"] == "native_product"
    assert any(m.get("answer_ref") == saved["answer_ref"] for m in refreshed["messages"])
    exported = export_assistant_reply(sid, answer["body"], "md")
    assert answer["body"] in (cfg.sessions_resolved / sid / "reports" / exported["artifact_path"].split("/")[-1]).read_text(encoding="utf-8")


def test_pending_multisource_replacement_only_requires_load_for_attempted_source(engine):
    import pandas as pd
    from data_agent.agent.intent import TurnIntent
    from data_agent.agent.publication.runtime import register_request
    from data_agent.tools.data_io import _record_trust_workflow
    from data_agent.tools.registry import registry
    loop, cfg = engine; ws = loop.context.workspace; state = loop.context.analysis_state
    ws.add("stable", pd.DataFrame({"amount": [5, 6]})); stable = ws.get_data_identity("stable")
    specs = [{"kind": "measurement", "scope": {"datasets": [n], "quantity": "rows"}} for n in ["orders", "stable"]]
    request(loop, "统计两表记录数", specs); old = state.answer_requirements[-1]["id"]
    content = "加载orders新版替换原目标，保留stable同版并统计两表记录数。"
    loop.messages.append({"role": "user", "content": content, "request_id": "multisource_replace"})
    intent = TurnIntent("directed_analysis", "clear", "data_loaded", "execute", "load_then_analyze", execution_readiness="pending_load", requested_outputs=[{"exact_span": content, "kind": "result", "op": "replace", "target_id": old, "business_scope": {"datasets": ["orders", "stable"], "quantity": "rows"}}])
    register_request(loop, content, intent)
    source = cfg.workspace_resolved / "inbox" / "orders_new.csv"; source.parent.mkdir(parents=True, exist_ok=True); source.write_text("amount\n99\n", encoding="utf-8")
    args = {"source": str(source), "name": "orders"}
    binding = loop._capture_tool_execution_binding("reload_data", args)
    loaded = registry.execute("reload_data", args)
    assert "error" not in loaded.summary.lower(), loaded.summary
    assert loop._record_turn_tool_result("reload_data", loaded.summary, args, "multi_load", execution_binding=binding)
    current = state.answer_requirements[-1]
    assert current["source_binding"]["required_load_datasets"] == ["orders"]
    for name in ["orders", "stable"]:
        _record_trust_workflow(session_id=loop.session_id, state=state, dataset=name, df=ws.get(name), applied=[], needs_confirm=[], quality={}, interpretation_data={}, detail_path="")
    declarations = [dict(d, requirement_id=current["id"], label=n) for d, n in zip(specs, ["orders", "stable"])]
    plan = {"contract_version": "stage3c0b.v2", "requirement_revision": state.requirement_revision, "goal": content, "method_plan": [{"goal": "Count independently", "expected_output": "Two row counts", "evidence_requirements": ["receipt"], "dataset_inputs": ["orders", "stable"], "combination_mode": "independent", "required_deliveries": declarations}]}
    result = registry.execute("record_analysis_plan", {"plan_json": plan})
    payload = result.data or json.loads(result.summary)
    assert "error" not in payload, payload
    current = state.answer_requirements[-1]
    assert current["source_binding"]["status"] == "frozen"
    assert current["source_data_identities"]["stable"] == stable == ws.get_data_identity("stable")
    assert current["source_data_identities"]["orders"]["fingerprint"] != state.answer_requirements[-2]["source_data_identities"]["orders"]["fingerprint"]


@pytest.mark.parametrize("case", ["select", "with", "empty", "collision"])
def test_normal_app_sql_snapshot_plan_delivery_and_failed_replacement(engine, tmp_path, case):
    import sqlalchemy
    from data_agent.web.app import create_app
    from data_agent.web.agent_manager import AgentManager
    from data_agent.session.history import load_session
    from data_agent.agent.publication.runtime import read_answer
    from data_agent.tools.report import export_assistant_reply
    from data_agent.llm.client import Response, ToolCall
    from data_agent.agent.context import use_agent_context
    _, cfg = engine
    database = tmp_path / "app.sqlite"
    connection = "sqlite:///" + str(database)
    db = sqlalchemy.create_engine(connection)
    with db.begin() as conn:
        conn.exec_driver_sql("CREATE TABLE incoming (amount INTEGER, code TEXT)")
        conn.exec_driver_sql("INSERT INTO incoming VALUES (7, '12%'), (13, '2026-10-11'), (22, 'raw')")
    db.dispose()
    query = "WITH observed AS (SELECT * FROM incoming) SELECT * FROM observed" if case == "with" else "SELECT * FROM incoming"
    if case == "empty": query += " WHERE 1=0"
    content = "从指定数据库载入查询结果，命名为experiment，报告本次查询实际记录数。"
    class SqlModel(_NativeProductModel):
        phase = 0
        loaded = False
        def chat(self, messages, system="", **kw):
            from scripts.testing.run_p7_browser_rehearsal import ScriptedClient
            if system == "你是一个JSON输出器，只返回JSON，不要其他内容。":
                output = {"kind": "result", "exact_span": self.content, "business_scope": {"datasets": ["experiment"], "quantity": "rows"}}
                if self.phase:
                    declared, _ = json.JSONDecoder().raw_decode(messages[-1]["content"].partition("canonical_requirements: ")[2])
                    prior = next(r for r in reversed(declared) if r["lifecycle_status"] == "active" and r["semantic_kind"] == "result")
                    output.update(op="replace", target_id=prior["id"])
                return Response(text=encode({"intent_type": "directed_analysis", "requested_outputs": [output]}))
            if not self.loaded:
                self.loaded = True; self.calls.append("load_sql")
                return Response(tool_calls=[ToolCall("sql_load_" + str(self.phase), "load_sql", {"name": "experiment", "connection_string": connection, "query": query})])
            if self.phase and self.step == 2:
                assert "replacement_source_unverified" in messages[-1]["content"]
                view = ScriptedClient._current_host_snapshot(messages, system)
                return Response(text=encode({"schema": "answer_bindings.v2", "snapshot_digest": view["snapshot_digest"], "sections": []}))
            if self.step == 0:
                self.step = 1; self.calls.append("quick_profile")
                return Response(tool_calls=[ToolCall("sql_profile", "quick_profile", {"name": "experiment"})])
            return super().chat(messages, system=system, **kw)
    specs = [{"kind": "measurement", "scope": {"datasets": ["experiment"], "quantity": "rows"}}]
    model = SqlModel(content, "experiment", "describe_dataset", {"name": "experiment"}, specs)
    manager = AgentManager(client_factory=lambda **kw: model, auxiliary_client_factory=lambda: model)
    app = create_app(); app.config["agent_manager"] = manager
    sid = "s3_sql_" + case
    def finish(response):
        assert response.status_code == 200
        events = [(b.split("\n", 1)[0][7:], json.loads(b.split("\ndata: ", 1)[1])) for b in response.get_data(as_text=True).split("\n\n") if b.startswith("event:")]
        ends = [d for e, d in events if e == "turn_end"]
        assert len(ends) == 1 and ends[0]["status"] == "completed", events
        saved = load_session(sid)["messages"][-1]
        answer = read_answer(sid, saved["answer_ref"])
        assert answer["body"] == saved["content"] == ends[0]["reply_content"]
        return saved, answer
    with app.test_client() as http:
        saved, answer = finish(http.post("/api/chat", json={"session_id": sid, "message": content}))
        assert answer["requirement_outcomes"][0]["status"] == "answered", answer
        active = next(r for r in answer["source_snapshot"]["requirements"] if not r["withdrawn"])
        row_claim = next(c for c in answer["source_snapshot"]["claims"] if c["id"] in active["claim_ids"])
        assert row_claim["value"] == (0 if case == "empty" else 3)
        with use_agent_context(manager.get(sid).context):
            ws = manager.get(sid).context.workspace
            original_identity = ws._publication_data_identity("experiment")
            origin = original_identity["source_origins"][0]
            assert origin["kind"] == "sql_query" and origin["query_sha256"] == "sha256:" + hashlib.sha256(query.encode()).hexdigest()
            import pandas as pd
            backup = cfg.sessions_resolved / sid / "data" / "experiment.parquet"
            frame = pd.read_parquet(backup) if backup.exists() else pd.read_pickle(backup.with_suffix(".pkl"))
            assert frame["code"].tolist() == ([] if case == "empty" else ["12%", "2026-10-11", "raw"])
            meta = json.loads((cfg.sessions_resolved / sid / "workspace_meta.json").read_text(encoding="utf-8"))
            assert meta["experiment"]["data_identity"] == original_identity
            parents = [d["data_identity"] for d in meta.values() if d["data_identity"]["version_id"] in original_identity["parent_version_ids"]]
            assert len(parents) == 1 and parents[0]["role"] == "raw" and parents[0]["source_origins"] == original_identity["source_origins"]
        assert model.calls == ["load_sql", "quick_profile", "record_analysis_plan", "describe_dataset"]
        receipt = next(r for r in manager.get(sid).context.analysis_state.tool_receipts if r["tool_name"] == "describe_dataset")
        assert receipt["plan_id"] and receipt["step_id"] == "native_product"
        if case == "collision":
            prior = deepcopy(answer)
            model.phase = 1; model.loaded = False; model.step = 0
            model.content = "明确从数据库加载新版并替换原目标，报告experiment新版实际记录数。"
            _, failed = finish(http.post("/api/chat", json={"session_id": sid, "message": model.content}))
            assert read_answer(sid, saved["answer_ref"]) == prior
            assert failed["requirement_outcomes"][-1]["status"] == "evidence_insufficient", failed
            with use_agent_context(manager.get(sid).context):
                assert ws._publication_data_identity("experiment") == original_identity
                assert not ws._publication_data_identity("experiment__history")
            assert model.calls[-3:] == ["load_sql", "quick_profile", "record_analysis_plan"]
            assert "新版来源尚未成功加载核验" in failed["body"]
            active_req = next(r for r in manager.get(sid).context.analysis_state.answer_requirements if r["lifecycle_status"] == "active")
            assert active_req["source_binding"]["status"] == "pending"
            assert active_req["source_binding"]["required_load_datasets"] == ["experiment"]
            assert not active_req["source_binding"].get("verified_load_receipts")
            answer = failed
        refreshed = http.get("/api/sessions/" + sid).get_json()
        assert any(m.get("answer_ref") for m in refreshed["messages"])
    exported = export_assistant_reply(sid, answer["body"], "md")
    assert exported["status"] == "exported"
    assert answer["body"] in (cfg.sessions_resolved / sid / "reports" / exported["artifact_path"].split("/")[-1]).read_text(encoding="utf-8")
    database.unlink()  # Restoring an accepted query snapshot never re-queries a live database.
    restored_manager = AgentManager(client_factory=lambda **kw: model, auxiliary_client_factory=lambda: model)
    restored = restored_manager.get_or_create(sid)
    with use_agent_context(restored.context):
        assert restored.context.workspace._publication_data_identity("experiment") == original_identity
    original_answer = read_answer(sid, saved["answer_ref"])
    assert original_answer["body"] == saved["content"]


@pytest.mark.parametrize("case", ["write", "unsupported"])
def test_normal_app_sql_read_only_refusal_has_exact_visible_reason(engine, tmp_path, case):
    import sqlite3
    from data_agent.web.app import create_app
    from data_agent.web.agent_manager import AgentManager
    from data_agent.session.history import load_session
    from data_agent.agent.publication.runtime import read_answer
    from data_agent.tools.report import export_assistant_reply
    from data_agent.llm.client import Response, ToolCall
    _, cfg = engine
    database = tmp_path / "refusal.sqlite"
    with sqlite3.connect(database) as conn:
        conn.execute("CREATE TABLE incoming(amount INTEGER)")
        conn.executemany("INSERT INTO incoming VALUES (?)", [(7,), (13,), (22,)])
    connection = "sqlite:///" + str(database) if case == "write" else "postgresql://user:secret@example.test/db"
    query = "WITH observed AS (SELECT 1) DELETE FROM incoming RETURNING amount" if case == "write" else "SELECT 1"
    content = "从指定数据库查询载入experiment，报告本次查询记录数。"
    class DeniedModel(_NativeProductModel):
        def chat(self, messages, system="", **kw):
            from scripts.testing.run_p7_browser_rehearsal import ScriptedClient
            if system == "你是一个JSON输出器，只返回JSON，不要其他内容。":
                return Response(text=encode({"intent_type": "directed_analysis", "requested_outputs": [{"kind": "result", "exact_span": content, "business_scope": {"datasets": ["experiment"], "quantity": "rows"}}]}))
            if self.step == 0:
                self.step += 1; self.calls.append("load_sql")
                return Response(tool_calls=[ToolCall("sql_refusal", "load_sql", {"name": "experiment", "connection_string": connection, "query": query})])
            view = ScriptedClient._current_host_snapshot(messages, system)
            return Response(text=encode({"schema": "answer_bindings.v2", "snapshot_digest": view["snapshot_digest"], "sections": []}))
    model = DeniedModel(content, "experiment", "", {}, [])
    manager = AgentManager(client_factory=lambda **kw: model, auxiliary_client_factory=lambda: model)
    app = create_app(); app.config["agent_manager"] = manager; sid = "s3_sql_refusal_" + case
    with app.test_client() as http:
        response = http.post("/api/chat", json={"session_id": sid, "message": content})
        events = [(b.split("\n", 1)[0][7:], json.loads(b.split("\ndata: ", 1)[1])) for b in response.get_data(as_text=True).split("\n\n") if b.startswith("event:")]
        ends = [d for e, d in events if e == "turn_end"]
        assert len(ends) == 1 and ends[0]["status"] == "completed", events
        saved = load_session(sid)["messages"][-1]; answer = read_answer(sid, saved["answer_ref"])
        assert answer["body"] == saved["content"] == ends[0]["reply_content"]
        assert answer["requirement_outcomes"][0]["status"] == "failed", answer
        assert "只读" in answer["body"] and ("拒绝数据变更" if case == "write" else "未支持") in answer["body"]
        assert "secret" not in answer["body"] and not answer["source_snapshot"]["claims"]
        assert model.calls == ["load_sql"]
        refreshed = http.get("/api/sessions/" + sid).get_json()
        assert any(m.get("answer_ref") == saved["answer_ref"] for m in refreshed["messages"])
    with sqlite3.connect(database) as conn:
        assert conn.execute("SELECT amount FROM incoming ORDER BY amount").fetchall() == [(7,), (13,), (22,)]
    exported = export_assistant_reply(sid, answer["body"], "md")
    assert answer["body"] in (cfg.sessions_resolved / sid / "reports" / exported["artifact_path"].split("/")[-1]).read_text(encoding="utf-8")


@pytest.mark.parametrize("other_failed", [False, True])
def test_sql_source_failure_clears_on_verified_retry_only_for_same_source(engine, other_failed):
    from data_agent.tools.registry import registry
    from data_agent.agent.publication.requirements import bind_requirements
    loop, _ = engine
    request(loop, "载入指定SQL来源并报告记录数")
    req = loop.context.analysis_state.answer_requirements[-1]
    req["business_scope"] = {"datasets": ["sql_retry"] + (["second"] if other_failed else [])}
    def run(name, connection, query):
        args = {"name": name, "connection_string": connection, "query": query}
        binding = loop._capture_tool_execution_binding("load_sql", args)
        result = registry.execute("load_sql", args)
        try: structured = result.data or json.loads(result.summary)
        except ValueError: structured = None
        return loop._record_turn_tool_result("load_sql", result.summary, args, "sql_" + name + "_" + str(len(loop._turn_tools_used)), structured, execution_binding=binding)
    assert not run("sql_retry", "postgresql://unused/db", "SELECT 1")
    assert bind_requirements(loop.context.analysis_state, [])[req["id"]]["gap"]["status"] == "failed"
    if other_failed: assert not run("second", "postgresql://unused/db", "SELECT 1")
    assert run("sql_retry", "sqlite://", "SELECT 7 AS amount")
    assert set(req["source_binding"]["load_failures"]) == ({"second"} if other_failed else set())
    gap = bind_requirements(loop.context.analysis_state, [])[req["id"]]["gap"]
    assert gap["status"] == ("failed" if other_failed else "not_computed")
    assert loop.context.workspace.get("sql_retry")["amount"].tolist() == [7]


@pytest.mark.parametrize("tool", ["export_data", "export_output"])
def test_json_export_exact_finite_time_missing_and_native_values(engine, tool):
    from pathlib import Path
    from datetime import date, datetime, time, timezone
    import pandas as pd
    from data_agent.tools.registry import registry
    loop, cfg = engine
    stamp = pd.Timestamp("2026-10-11T12:34:56.123456789+08:00")
    frame = pd.DataFrame({"value": [0.423456789012345, 1.2345678901234567, 5e-324],
        "integer": pd.Series([9007199254740993, -9007199254740993, None], dtype="Int64"),
        "time": [stamp, pd.NaT, stamp], "flag": pd.Series([True, False, None], dtype="boolean"),
        "text": ["中文", "", None], "object": [date(2026, 10, 11), datetime(2026, 10, 11, 12, 34, 56, 123456, tzinfo=timezone.utc), time(12, 34, 56, 123456, tzinfo=timezone.utc)]})
    loop.context.workspace.add("exact", frame)
    args = {"name": "exact", "path": str(cfg.workspace_resolved / "data/exact.json"), "fmt": "json"}
    if tool == "export_output": args["output_type"] = "data"
    result = registry.execute(tool, args); data = result.data or json.loads(result.summary)
    assert "error" not in data, data
    expected = [{"value": float(frame.iloc[i]["value"]), "integer": [9007199254740993, -9007199254740993, None][i],
        "time": None if i == 1 else stamp.isoformat(), "flag": [True, False, None][i], "text": ["中文", "", None][i],
        "object": ["2026-10-11", "2026-10-11T12:34:56.123456+00:00", "12:34:56.123456+00:00"][i]} for i in range(3)]
    original = Path(data["destination"]).read_bytes()
    assert json.loads(original) == expected
    assert (cfg.sessions_resolved.parent / data["artifact_path"]).read_bytes() == original
    assert data["sha256"] == "sha256:" + hashlib.sha256(original).hexdigest()
    assert frame.equals(loop.context.workspace.get("exact"))


@pytest.mark.parametrize("tool", ["export_data", "export_output"])
@pytest.mark.parametrize("case", ["positive_inf", "negative_inf", "decimal", "complex", "bytes", "column_type", "duplicate_columns"])
def test_json_export_refuses_lossy_values_before_file_or_download(engine, tool, case):
    from decimal import Decimal
    from pathlib import Path
    import pandas as pd
    from data_agent.tools.registry import registry
    from data_agent.session.history import load_session
    loop, cfg = engine
    values = {"positive_inf": float("inf"), "negative_inf": -float("inf"), "decimal": Decimal("1.234567890123456789"), "complex": 1+2j, "bytes": b"abc"}
    frame = pd.DataFrame({"value": [values.get(case, 1)]})
    if case == "column_type": frame.columns = [1]
    if case == "duplicate_columns": frame = pd.DataFrame([[1, 2]], columns=["value", "value"])
    loop.context.workspace.add("lossy", frame)
    destination = cfg.workspace_resolved / "data/existing.json"; destination.parent.mkdir(parents=True, exist_ok=True)
    destination.write_bytes(b"old immutable working contents")
    before = deepcopy(load_session(loop.session_id))
    args = {"name": "lossy", "path": str(destination), "fmt": "json"}
    if tool == "export_output": args["output_type"] = "data"
    result = registry.execute(tool, args); data = result.data or json.loads(result.summary)
    assert "error" in data and "JSON" in data["error"], data
    assert destination.read_bytes() == b"old immutable working contents"
    assert not list((cfg.sessions_resolved / loop.session_id / "output").glob("export_*"))
    assert load_session(loop.session_id) == before


@pytest.mark.parametrize('tool,case', [('export_data', 'csv'), ('export_data', 'excel'), ('export_data', 'json'), ('export_output', 'csv'), ('export_output', 'excel'), ('export_output', 'json'), ('export_data', 'positive_inf'), ('export_output', 'negative_inf'), ('export_data', 'excel_high'), ('export_output', 'excel_time'), ('export_data', 'excel_bytes'), ('export_output', 'excel_bytes')])
def test_normal_app_export_alias_format_exact_download_sse_refresh_and_saved_identity(engine, tool, case):
    from pathlib import Path
    import pandas as pd
    from data_agent.web.app import create_app
    from data_agent.web.agent_manager import AgentManager
    from data_agent.session.history import load_session
    from data_agent.agent.publication.runtime import read_answer
    from data_agent.tools.report import export_assistant_reply
    from data_agent.agent.context import use_agent_context
    from data_agent.llm.client import Response
    _, cfg = engine
    denied = case in {"positive_inf", "negative_inf", "excel_bytes"}
    fmt = "json" if case in {"positive_inf","negative_inf"} else "excel" if case.startswith("excel_") else case
    content = "只将当前experiment完整导出为指定格式文件，提供本次下载。"
    args = {"name": "experiment", "path": "output/current." + {"excel": "xlsx"}.get(fmt, fmt), "fmt": fmt}
    if tool == "export_output": args["output_type"] = "data"
    specs = [{"kind": "artifact", "scope": {"datasets": ["experiment"]}}]
    class ExportModel(_NativeProductModel):
        def chat(self, messages, system="", **kw):
            from scripts.testing.run_p7_browser_rehearsal import ScriptedClient
            if system == "你是一个JSON输出器，只返回JSON，不要其他内容。":
                return Response(text=encode({"intent_type": "data_operation", "requested_outputs": [{"kind": "result", "exact_span": content, "business_scope": {"datasets": ["experiment"]}}]}))
            if self.step == 0:
                from data_agent.llm.client import ToolCall
                self.step = 1; self.calls.append("quick_profile")
                return Response(tool_calls=[ToolCall("export_profile", "quick_profile", {"name": "experiment"})])
            if self.step < 3: return super().chat(messages, system=system, **kw)
            view = ScriptedClient._current_host_snapshot(messages, system)
            return Response(text=encode({"schema": "answer_bindings.v2", "snapshot_digest": view["snapshot_digest"], "sections": [{"requirement_ids": [r["id"]], "parts": [{"kind": "artifact", "ref": a} for a in r.get("artifact_ids", [])]} for r in view["requirements"] if not r["withdrawn"]]}))
    model = ExportModel(content, "experiment", tool, args, specs)
    manager = AgentManager(client_factory=lambda **kw: model, auxiliary_client_factory=lambda: model)
    sid = "s3_export_" + tool + "_" + case
    loop = manager.get_or_create(sid)
    frame = pd.DataFrame({"amount": [0.423456789012345, 2.5], "text": ["中文", "original"]})
    if fmt == "json":
        frame["amount"] = [1.2345678901234567, 5e-324]
        frame["when"] = [pd.Timestamp("2026-10-11T12:34:56.123456789+08:00"), pd.NaT]
    if case == "excel_high":
        frame["amount"] = [1.2345678901234567, 2.5]
        frame["large"] = pd.Series([9007199254740993, 42], dtype="Int64")
        frame["text"] = ["=1+1", "plain"]
    if case == "excel_time":
        frame["when"] = [pd.Timestamp("2026-10-11T12:34:56.123456789+08:00"), pd.NaT]
        frame["elapsed"] = [pd.Timedelta("1 days 00:00:00.123456789"), pd.NaT]
    if case == "excel_bytes": frame["raw_bytes"] = [b"abc", b"def"]
    if case in {"positive_inf","negative_inf"}: frame["amount"] = [float("inf") if case == "positive_inf" else -float("inf"), 2.5]
    with use_agent_context(loop.context):
        loop.context.workspace.add("experiment", frame)
        identity = loop.context.workspace._publication_data_identity("experiment")
        from data_agent.tools.data_io import _record_trust_workflow
        _record_trust_workflow(session_id=sid, state=loop.context.analysis_state, dataset="experiment", df=frame, applied=[], needs_confirm=[], quality={}, interpretation_data={}, detail_path="")
    app = create_app(); app.config["agent_manager"] = manager
    with app.test_client() as http:
        response = http.post("/api/chat", json={"session_id": sid, "message": content})
        events = [(b.split("\n", 1)[0][7:], json.loads(b.split("\ndata: ", 1)[1])) for b in response.get_data(as_text=True).split("\n\n") if b.startswith("event:")]
        ends = [d for e, d in events if e == "turn_end"]
        assert len(ends) == 1 and ends[0]["status"] == "completed", events
        saved = load_session(sid)["messages"][-1]; answer = read_answer(sid, saved["answer_ref"])
        assert answer["body"] == saved["content"] == ends[0]["reply_content"]
        if denied:
            assert answer["requirement_outcomes"][0]["status"] != "answered", answer
            assert not answer["source_snapshot"]["artifacts"] and not answer["artifact_refs"]
            assert not list((cfg.sessions_resolved / sid / "output").glob("export_*"))
            assert any(m.get("role") == "tool" and ("Excel导出拒绝" if case=="excel_bytes" else "JSON导出拒绝") in m.get("content", "") for m in load_session(sid)["messages"])
            assert model.calls == ["quick_profile", "record_analysis_plan", tool]
            assert not any(r["tool_name"] == tool for r in loop.context.analysis_state.tool_receipts)
            refreshed = http.get("/api/sessions/" + sid).get_json()
            assert any(m.get("answer_ref") == saved["answer_ref"] for m in refreshed["messages"])
            exported = export_assistant_reply(sid, answer["body"], "md")
            assert answer["body"] in (cfg.sessions_resolved / sid / "reports" / exported["artifact_path"].split("/")[-1]).read_text(encoding="utf-8")
            return
        assert answer["requirement_outcomes"][0]["status"] == "answered", answer
        req = next(r for r in answer["source_snapshot"]["requirements"] if not r["withdrawn"])
        assert req["artifact_ids"] and not req["claim_ids"] and not req["effect_ids"]
        artifact = next(a for a in answer["source_snapshot"]["artifacts"] if a["id"] == req["artifact_ids"][0])
        download = http.get("/api/files/sessions/" + sid + "/" + artifact["path"])
        assert download.status_code == 200, artifact
        actual = cfg.sessions_resolved / sid / artifact["path"]
        assert download.data == actual.read_bytes()
        if fmt == "json":
            records = json.loads(download.data)
            assert [r["amount"] for r in records] == frame["amount"].tolist()
            assert [r["when"] for r in records] == [frame.iloc[0]["when"].isoformat(), None]
        elif case == "excel_time":
            from openpyxl import load_workbook
            workbook = load_workbook(actual,data_only=False);sheet=workbook.active
            assert sheet["C2"].value==frame["when"].iloc[0].isoformat() and sheet["C2"].data_type=="s"
            assert sheet["D2"].value==frame["elapsed"].iloc[0].isoformat() and sheet["D2"].data_type=="s"
            assert sheet["C3"].value is None and sheet["D3"].value is None
            assert "时间或时长转文本" in answer["body"] and "when" in answer["body"] and "elapsed" in answer["body"]
            workbook.close()
        elif case == "excel_high":
            from openpyxl import load_workbook
            workbook = load_workbook(actual, data_only=False); sheet = workbook.active
            assert sheet['A2'].value == '1.2345678901234567' and sheet['A2'].data_type == 's'
            assert sheet['C2'].value == '9007199254740993' and sheet['C2'].data_type == 's'
            assert sheet['B2'].value == '=1+1' and sheet['B2'].data_type == 's'
            assert sheet['A3'].value == 2.5 and sheet['A3'].data_type == 'n'
            assert '数值转文本' in answer['body'] and 'amount' in answer['body'] and 'large' in answer['body']
            workbook.close()
        else:
            restored = pd.read_csv(actual) if fmt == "csv" else pd.read_excel(actual)
            assert restored["amount"].tolist() == frame["amount"].tolist()
            assert restored["text"].tolist() == frame["text"].tolist()
        refreshed = http.get("/api/sessions/" + sid).get_json()
        assert any(m.get("answer_ref") == saved["answer_ref"] for m in refreshed["messages"])
        working = cfg.sessions_resolved / sid / "output" / ("current." + {"excel": "xlsx"}.get(fmt, fmt))
        working.write_bytes(b"later working file overwrite")
        assert http.get("/api/files/sessions/" + sid + "/" + artifact["path"]).data == download.data
    assert model.calls == ["quick_profile", "record_analysis_plan", tool]
    receipt = next(r for r in loop.context.analysis_state.tool_receipts if r["tool_name"] == tool)
    assert receipt["plan_id"] and receipt["step_id"] == "native_product"
    assert receipt["data_identities"]["experiment"] == identity
    if case == "excel_high":
        from data_agent.agent.publication.facts import load_receipt_result
        with use_agent_context(loop.context):
            _, _, raw, _ = load_receipt_result(session_id=sid, receipt_id=receipt['id'])
        assert raw['format_audit']['numeric_as_text_total'] == 2
    exported = export_assistant_reply(sid, answer["body"], "md")
    assert answer["body"] in (cfg.sessions_resolved / sid / "reports" / exported["artifact_path"].split("/")[-1]).read_text(encoding="utf-8")


@pytest.mark.parametrize("tool", ["export_data", "export_output"])
def test_excel_export_numeric_text_mapping_and_literal_strings(engine, tool):
    from pathlib import Path
    from decimal import Decimal
    import pandas as pd
    from openpyxl import load_workbook
    from data_agent.tools.registry import registry
    loop, cfg = engine
    frame = pd.DataFrame({"amount": [1.2345678901234567, 2.5, -0.0],
        "large": pd.Series([9007199254740993, 42, None], dtype="Int64"),
        "text": ["=1+1", "plain", "中文"], "decimal": [Decimal("1.234567890123456789"), Decimal("2.5"), None]})
    loop.context.workspace.add("excel_exact", frame)
    args = {"name": "excel_exact", "path": str(cfg.workspace_resolved / "data/exact.xlsx"), "fmt": "excel"}
    if tool == "export_output": args["output_type"] = "data"
    result = registry.execute(tool, args); data = result.data or json.loads(result.summary)
    assert "error" not in data, data
    workbook = load_workbook(data["destination"], data_only=False); sheet = workbook.active
    assert [(sheet[f"A{i}"].value, sheet[f"A{i}"].data_type) for i in range(2, 5)] == [("1.2345678901234567", "s"), (2.5, "n"), ("-0.0", "s")]
    assert [(sheet[f"B{i}"].value, sheet[f"B{i}"].data_type) for i in range(2, 4)] == [("9007199254740993", "s"), (42, "n")]
    assert sheet['C2'].value == '=1+1' and sheet['C2'].data_type == 's'
    assert sheet['D2'].value == '1.234567890123456789' and sheet['D2'].data_type == 's'
    assert sheet['D3'].value == 2.5 and sheet['D3'].data_type == 'n'
    audit = data['format_audit']; assert audit['numeric_as_text_total'] == 4
    assert [(c['column'], c['numeric_as_text_count'], c['literal_string_count']) for c in audit['columns']] == [('amount', 2, 0), ('large', 1, 0), ('text', 0, 3), ('decimal', 1, 0)]
    assert '不能直接' in audit['limitation'] and '数值转文本' in data['message']
    workbook.close()
    cached = load_workbook(data['destination'], data_only=True);assert cached.active['C2'].value == '=1+1';cached.close()
    original = Path(data['destination']).read_bytes()
    assert (cfg.sessions_resolved.parent / data['artifact_path']).read_bytes() == original
    assert data['sha256'] == 'sha256:' + hashlib.sha256(original).hexdigest()
    assert loop.context.workspace.get('excel_exact').equals(frame)


@pytest.mark.parametrize("tool", ["export_data", "export_output"])
def test_excel_export_refuses_silent_text_truncation_before_overwrite(engine, tool):
    import pandas as pd
    from data_agent.tools.registry import registry
    loop, cfg = engine
    loop.context.workspace.add('long_text', pd.DataFrame({'text': ['x' * 33000]}))
    destination = cfg.workspace_resolved / 'data/old.xlsx';destination.parent.mkdir(parents=True, exist_ok=True);destination.write_bytes(b'old contents')
    args = {'name': 'long_text', 'path': str(destination), 'fmt': 'excel'}
    if tool == 'export_output': args['output_type'] = 'data'
    result = registry.execute(tool, args);data = result.data or json.loads(result.summary)
    assert 'error' in data and 'Excel' in data['error'], data
    assert destination.read_bytes() == b'old contents'
    assert not list((cfg.sessions_resolved / loop.session_id / 'output').glob('export_*'))


def _list_catalog_fixture(loop):
    import pandas as pd
    frame = pd.DataFrame({f"c_{i:02d}": [i, i + 1] for i in range(12)})
    loop.context.workspace.add("catalog", frame)
    for i in range(12):
        loop.context.workspace.derive("catalog", f"derived_{i:02d}", frame, "fixture transformation " + str(i))
    return frame


def test_list_data_full_columns_history_and_receipt_scope(engine):
    from data_agent.tools.registry import registry
    loop, _ = engine
    frame = _list_catalog_fixture(loop)
    rid, result = execute(loop, "list_data", {})
    assert isinstance(result, dict), result
    assert result["datasets"]["catalog"]["column_names"] == list(frame.columns)
    assert len(result["transform_history"]) == 12
    assert result["transform_history"][0] == {"from": "catalog", "op": "derive", "to": "derived_00"}
    assert result["transform_history"][-1]["to"] == "derived_11"
    receipt = loop.context.analysis_state.tool_receipts[-1]
    assert set(receipt["dataset_refs"]) == set(result["datasets"]) == set(receipt["data_identities"])
    from data_agent.agent.publication.native import resolve_native_claims
    facts = resolve_native_claims(session_id=loop.session_id, receipt_id=rid)
    assert {tuple(c["basis_refs"][0]["result_path"]) for c in facts} == {("datasets",), ("transform_history",)}
    assert all(c["target"]["business_product"]["kind"] == "table" for c in facts)
    assert facts[0]["basis_refs"][0]["data_identities"] == receipt["data_identities"]


@pytest.mark.parametrize("phase", ["planning", "execution", "synthesis", "error"])
def test_list_data_keeps_existing_visibility_and_never_leaks_hidden_parent(engine, phase):
    import pandas as pd
    from data_agent.tools.registry import registry
    from data_agent.agent.execution_scope import WorkspaceScopeSnapshot
    loop, _ = engine
    _list_catalog_fixture(loop)
    loop.context.workspace.add("secret", pd.DataFrame({"secret_column": [999]}))
    scope = WorkspaceScopeSnapshot(phase=phase, session_id=loop.session_id, allowed_datasets=frozenset({"derived_11"}))
    with loop.context.bind_workspace_scope(scope):
        result = registry.execute("list_data", {}).data
    assert isinstance(result, dict), result
    expected = {"derived_11"} if phase in {"planning", "execution"} else set()
    assert set(result["datasets"]) == expected
    assert result["transform_history"] == []
    assert "secret" not in encode(result) and "catalog" not in encode(result)
    if expected:
        assert result["datasets"]["derived_11"]["derived_from"] is None
        assert "data_identity" not in result["datasets"]["derived_11"] and "metadata" not in result["datasets"]["derived_11"]


@pytest.mark.parametrize("fault", ["missing_column", "missing_history", "invented_rows", "extra_dataset", "old_version", "foreign_session"])
def test_list_data_receipt_rejects_incomplete_forged_or_stale_catalog(engine, fault):
    import pandas as pd
    from data_agent.tools.registry import registry
    from data_agent.agent.publication.native import resolve_native_claims
    loop, _ = engine
    _list_catalog_fixture(loop)
    binding = loop._capture_tool_execution_binding("list_data", {})
    result = registry.execute("list_data", {}).data
    assert isinstance(result, dict)
    if fault == "missing_column": result["datasets"]["catalog"]["column_names"].pop()
    elif fault == "missing_history": result["transform_history"].pop(0)
    elif fault == "invented_rows": result["datasets"]["catalog"]["rows"] = 999
    elif fault == "extra_dataset": result["datasets"]["invented"] = deepcopy(result["datasets"]["catalog"])
    rid = loop._record_turn_tool_result("list_data", encode(result), {}, "catalog_tamper", result, execution_binding=binding)
    if fault == "old_version": loop.context.workspace.add("catalog", pd.DataFrame({"changed": [1]}))
    with pytest.raises((ValueError, OSError)):
        resolve_native_claims(session_id="foreign" if fault == "foreign_session" else loop.session_id, receipt_id=rid)


def test_list_data_existing_table_delivery_is_complete_and_explicit(engine):
    from data_agent.agent.publication.runtime import snapshot_for_turn, publish
    loop, _ = engine
    _list_catalog_fixture(loop)
    names = sorted(loop.context.workspace.list_datasets())
    specs = [{"kind": "table", "scope": {"datasets": names, "parameters": {"catalog_output": key}}} for key in ("datasets", "transform_history")]
    request(loop, "列出当前授权数据集的全部字段和完整变换历史", specs)
    execute(loop, "list_data", {})
    snapshot = snapshot_for_turn(loop)
    body, error = publish(loop, draft(snapshot))
    assert error is None and "c_11" in body and "derived_00" in body
    assert loop._accepted_answer["requirement_outcomes"][0]["status"] == "answered"
    assert len(snapshot.as_dict()["requirements"][0]["claim_ids"]) == 2


@pytest.mark.parametrize("scoped", [False, True])
def test_normal_app_list_data_plan_sse_saved_refresh_and_md_export(engine, scoped):
    import pandas as pd
    from data_agent.web.app import create_app
    from data_agent.web.agent_manager import AgentManager
    from data_agent.session.history import load_session
    from data_agent.agent.publication.runtime import read_answer
    from data_agent.tools.report import export_assistant_reply
    from data_agent.agent.context import use_agent_context
    from data_agent.llm.client import Response, ToolCall
    from scripts.testing.run_p7_browser_rehearsal import ScriptedClient
    _, cfg = engine
    names = ["catalog"] if scoped else ["catalog"] + [f"derived_{i:02d}" for i in range(12)]
    content = "列出指定数据集的全部字段和完整变换历史：" + ",".join(names)
    scope = {"datasets": names}
    specs = [{"kind": "table", "scope": dict(scope, parameters={"catalog_output": key})} for key in ("datasets", "transform_history")]
    class CatalogModel(_NativeProductModel):
        def chat(self, messages, system="", **kw):
            if system == "你是一个JSON输出器，只返回JSON，不要其他内容。":
                return Response(text=encode({"intent_type": "directed_analysis", "requested_outputs": [{"kind": "result", "exact_span": content, "business_scope": scope}]}))
            view = ScriptedClient._current_host_snapshot(messages, system)
            self.step += 1
            if self.step == 1:
                goal = next(r for r in view["execution_targets"] if r["semantic_kind"] == "result")
                declarations = [dict(d, requirement_id=goal["id"], label="完整目录 " + str(i)) for i, d in enumerate(specs)]
                plan = {"contract_version": "stage3c0b.v2", "requirement_revision": view["requirement_revision"], "goal": content, "visualization_strategy": [], "method_plan": [{"step_id": "catalog_delivery", "goal": content, "expected_output": "完整可见目录和历史", "dataset_inputs": names, "combination_mode": "independent", "evidence_requirements": ["original receipt"], "required_deliveries": declarations}]}
                self.calls.append("record_analysis_plan")
                return Response(tool_calls=[ToolCall("catalog_plan", "record_analysis_plan", {"plan_json": plan})])
            if self.step == 2:
                self.calls.append("list_data")
                return Response(tool_calls=[ToolCall("catalog_execute", "list_data", {})])
            return Response(text=encode({"schema": "answer_bindings.v2", "snapshot_digest": view["snapshot_digest"], "sections": [{"requirement_ids": [r["id"]], "parts": [{"kind": "claim", "ref": cid} for cid in r["claim_ids"]]} for r in view["requirements"] if not r["withdrawn"]]}))
    model = CatalogModel(content, None, "list_data", {}, specs)
    manager = AgentManager(client_factory=lambda **kw: model, auxiliary_client_factory=lambda: model)
    sid = "s3_list_" + str(scoped)
    loop = manager.get_or_create(sid)
    with use_agent_context(loop.context):
        frame = _list_catalog_fixture(loop)
        loop.context.workspace.add("secret", pd.DataFrame({"hidden_marker": [999]}))
        from data_agent.tools.data_io import _record_trust_workflow
        for name in names:
            _record_trust_workflow(session_id=sid, state=loop.context.analysis_state, dataset=name, df=frame, applied=[], needs_confirm=[], quality={}, interpretation_data={}, detail_path="")
    app = create_app(); app.config["agent_manager"] = manager
    with app.test_client() as http:
        response = http.post("/api/chat", json={"session_id": sid, "message": content})
        events = [(b.split("\n", 1)[0][7:], json.loads(b.split("\ndata: ", 1)[1])) for b in response.get_data(as_text=True).split("\n\n") if b.startswith("event:")]
        ends = [d for e, d in events if e == "turn_end"]
        assert len(ends) == 1 and ends[0]["status"] == "completed", events
        saved = load_session(sid)["messages"][-1]; answer = read_answer(sid, saved["answer_ref"])
        assert answer["body"] == saved["content"] == ends[0]["reply_content"]
        assert answer["requirement_outcomes"][0]["status"] == "answered", answer
        assert "c_11" in answer["body"] and "hidden_marker" not in answer["body"] and "secret" not in answer["body"]
        req = next(r for r in answer["source_snapshot"]["requirements"] if not r["withdrawn"])
        facts = {c["basis_refs"][0]["result_path"][0]: c["value"] for c in answer["source_snapshot"]["claims"] if c["id"] in req["claim_ids"]}
        # Scoped execution may include verified derived descendants of catalog;
        # those existing permissions must be retained, never unrelated secret.
        assert set(facts["datasets"]) == {"catalog", *[f"derived_{i:02d}" for i in range(12)]}
        assert len(facts["transform_history"]) == 12 and facts["transform_history"][0]["to"] == "derived_00"
        refreshed = http.get("/api/sessions/" + sid).get_json()
        assert any(m.get("answer_ref") == saved["answer_ref"] for m in refreshed["messages"])
    assert model.calls == ["record_analysis_plan", "list_data"]
    receipt = next(r for r in loop.context.analysis_state.tool_receipts if r["tool_name"] == "list_data")
    assert receipt["step_id"] == "catalog_delivery" and receipt["plan_id"] and len(receipt["required_delivery_ids"]) == 2
    assert set(receipt["data_identities"]) == set(facts["datasets"])
    exported = export_assistant_reply(sid, answer["body"], "md")
    assert answer["body"] in (cfg.sessions_resolved / sid / "reports" / exported["artifact_path"].split("/")[-1]).read_text(encoding="utf-8")


def test_list_data_empty_visible_scope_does_not_assert_empty_session(engine):
    from data_agent.tools.registry import registry
    from data_agent.agent.execution_scope import WorkspaceScopeSnapshot
    loop, _ = engine
    assert loop.context.workspace.exists("orders")
    with loop.context.bind_workspace_scope(WorkspaceScopeSnapshot(phase="execution", session_id=loop.session_id)):
        result = registry.execute("list_data", {}).data
    assert result == {"scope": "current_visible_datasets", "datasets": {}, "transform_history": []}


def test_list_data_same_alias_in_other_session_cannot_leak_or_reuse_catalog(engine):
    import pandas as pd
    from data_agent.agent.context import use_agent_context
    from data_agent.agent.loop import AgentLoop
    from data_agent.tools.registry import registry
    from data_agent.agent.publication.native import resolve_native_claims
    loop, _ = engine
    _list_catalog_fixture(loop)
    rid, before = execute(loop, "list_data", {})
    other = AgentLoop(client=object(), session_id="catalog_other_session")
    with use_agent_context(other.context):
        other.context.workspace.add("catalog", pd.DataFrame({"foreign_marker": [123]}))
        current = registry.execute("list_data", {}).data
        assert set(current["datasets"]) == {"catalog"} and current["datasets"]["catalog"]["column_names"] == ["foreign_marker"]
        with pytest.raises(ValueError):
            resolve_native_claims(session_id=other.session_id, receipt_id=rid)
    assert registry.execute("list_data", {}).data == before


def test_list_data_host_rechecks_changed_history_without_changing_frame_version(engine):
    from data_agent.agent.publication.native import resolve_native_claims
    loop, _ = engine
    _list_catalog_fixture(loop)
    rid, result = execute(loop, "list_data", {})
    loop.context.workspace.log_transform("catalog", "later_operation", "derived_00")
    with pytest.raises(ValueError, match="目录"):
        resolve_native_claims(session_id=loop.session_id, receipt_id=rid)


def _catalog_metadata_specs():
    return [{"kind":"table","scope":{"datasets":[],"parameters":{"catalog_scope":"current_visible","catalog_output":key}}} for key in ("datasets","transform_history")]


def test_catalog_metadata_empty_formal_delivery_uses_session_scope(engine):
    from data_agent.agent.publication.runtime import snapshot_for_turn,publish,commit,read_answer
    from data_agent.agent.execution_scope import WorkspaceScopeSnapshot
    from data_agent.session.history import load_session
    loop,_=engine
    request(loop,"列出当前授权可见数据",_catalog_metadata_specs()[0])
    with loop.context.bind_workspace_scope(WorkspaceScopeSnapshot(phase="execution",session_id=loop.session_id)):
        rid,result=execute(loop,"list_data",{})
    assert result["datasets"]=={} and loop.context.workspace.exists("orders")
    receipt=loop.context.analysis_state.tool_receipts[-1]
    assert receipt["catalog_binding"]["session_id"]==loop.session_id
    assert receipt["catalog_binding"]["scope_fingerprint"]==receipt["scope_fingerprint"]
    snapshot=snapshot_for_turn(loop)
    body,error=publish(loop,draft(snapshot))
    assert error is None and "单位未知" not in body and "授权可见" in body
    assert loop._accepted_answer["requirement_outcomes"][0]["status"]=="answered"
    assert snapshot.as_dict()["data_identities"]=={}
    assert all(c["scope"]["dataset"] is None and c["basis_refs"][0]["catalog_binding"] for c in snapshot.as_dict()["claims"])
    assert commit(loop,{"role":"assistant"},body) is None
    saved=load_session(loop.session_id)["messages"][-1]
    assert read_answer(loop.session_id,saved["answer_ref"])["body"]==body


@pytest.mark.parametrize("kind",["measurement","forecast","chart","fit","scenario"])
def test_catalog_metadata_cannot_replace_normal_data_identity(kind):
    from data_agent.agent.analysis_plan_contracts import delivery_schema
    from jsonschema import Draft7Validator
    scope=_catalog_metadata_specs()[0]["scope"]
    if kind=="measurement":scope["quantity"]="rows"
    errors=list(Draft7Validator(delivery_schema()).iter_errors({"requirement_id":"rq","kind":kind,"label":"fake","scope":scope}))
    assert errors


@pytest.mark.parametrize("fault",["scope","session","digest","missing_binding","added_visible","history"])
def test_catalog_metadata_rejects_tampered_scope_or_changed_namespace(engine,fault):
    import pandas as pd
    from data_agent.agent.publication.native import resolve_native_claims
    loop,_=engine
    rid,_=execute(loop,"list_data",{})
    receipt=loop.context.analysis_state.tool_receipts[-1]
    if fault=="scope":receipt["catalog_binding"]["scope"]["allowed_datasets"]=["invented"]
    elif fault=="session":receipt["catalog_binding"]["session_id"]="foreign"
    elif fault=="digest":receipt["catalog_binding"]["snapshot_sha256"]="sha256:"+"0"*64
    elif fault=="missing_binding":receipt.pop("catalog_binding",None)
    elif fault=="added_visible":loop.context.workspace.add("new_visible",pd.DataFrame({"x":[1]}))
    else:loop.context.workspace.log_transform("orders","new_history","orders")
    with pytest.raises((ValueError,OSError)):
        resolve_native_claims(session_id=loop.session_id,receipt_id=rid)


def test_catalog_metadata_commit_rechecks_namespace_after_render(engine):
    import pandas as pd
    from data_agent.agent.publication.runtime import snapshot_for_turn,publish,commit
    loop,_=engine
    request(loop,"列出当前授权可见目录",_catalog_metadata_specs()[0])
    execute(loop,"list_data",{})
    snapshot=snapshot_for_turn(loop)
    body,error=publish(loop,draft(snapshot))
    assert error is None and loop._accepted_answer["requirement_outcomes"][0]["status"]=="answered"
    loop.context.workspace.add("late_visible",pd.DataFrame({"x":[1]}))
    assert commit(loop,{"role":"assistant"},body)
    assert loop.context.analysis_state.accepted_answers==[]


def test_normal_app_catalog_metadata_empty_plan_sse_refresh_export(engine):
    from data_agent.web.app import create_app
    from data_agent.web.agent_manager import AgentManager
    from data_agent.session.history import load_session
    from data_agent.agent.publication.runtime import read_answer
    from data_agent.tools.report import export_assistant_reply
    from data_agent.llm.client import Response,ToolCall
    from scripts.testing.run_p7_browser_rehearsal import ScriptedClient
    _,cfg=engine
    content="列出当前授权可见的数据集和全部变换历史"
    specs=_catalog_metadata_specs()
    class EmptyCatalogModel(_NativeProductModel):
        def chat(self,messages,system="",**kw):
            if system=="你是一个JSON输出器，只返回JSON，不要其他内容。":
                return Response(text=encode({"intent_type":"directed_analysis","requested_outputs":[{"kind":"result","exact_span":content,"business_scope":{"datasets":[],"parameters":{"catalog_scope":"current_visible"}}}]}))
            view=ScriptedClient._current_host_snapshot(messages,system)
            self.step+=1
            if self.step==1:
                goal=next(r for r in view["execution_targets"] if r["semantic_kind"]=="result")
                declarations=[dict(d,requirement_id=goal["id"],label="目录元数据 "+str(i)) for i,d in enumerate(specs)]
                plan={"contract_version":"stage3c0b.v2","requirement_revision":view["requirement_revision"],"goal":content,"visualization_strategy":[],"method_plan":[{"step_id":"empty_catalog","goal":content,"expected_output":"真实授权目录","dataset_inputs":[],"combination_mode":"independent","evidence_requirements":["original receipt with scope"],"required_deliveries":declarations}]}
                self.calls.append("record_analysis_plan")
                return Response(tool_calls=[ToolCall("plan","record_analysis_plan",{"plan_json":plan})])
            if self.step==2:
                self.calls.append("list_data")
                return Response(tool_calls=[ToolCall("catalog","list_data",{})])
            return Response(text=encode({"schema":"answer_bindings.v2","snapshot_digest":view["snapshot_digest"],"sections":[{"requirement_ids":[r["id"]],"parts":[{"kind":"claim","ref":cid} for cid in r["claim_ids"]]} for r in view["requirements"] if not r["withdrawn"]]}))
    model=EmptyCatalogModel(content,None,"list_data",{},specs)
    manager=AgentManager(client_factory=lambda **kw:model,auxiliary_client_factory=lambda:model)
    sid="s3_empty_catalog_metadata"
    loop=manager.get_or_create(sid)
    app=create_app();app.config["agent_manager"]=manager
    with app.test_client() as http:
        response=http.post("/api/chat",json={"session_id":sid,"message":content})
        events=[(b.split("\n",1)[0][7:],json.loads(b.split("\ndata: ",1)[1])) for b in response.get_data(as_text=True).split("\n\n") if b.startswith("event:")]
        ends=[d for e,d in events if e=="turn_end"]
        assert len(ends)==1 and ends[0]["status"]=="completed",events
        saved=load_session(sid)["messages"][-1];answer=read_answer(sid,saved["answer_ref"])
        assert answer["body"]==saved["content"]==ends[0]["reply_content"]
        assert answer["requirement_outcomes"][0]["status"]=="answered",answer
        facts={c["basis_refs"][0]["result_path"][0]:c["value"] for c in answer["source_snapshot"]["claims"]}
        assert facts=={"datasets":{},"transform_history":[]} and answer["source_snapshot"]["data_identities"]=={}
        assert "单位未知" not in answer["body"] and "授权可见" in answer["body"]
        assert any(m.get("answer_ref")==saved["answer_ref"] for m in http.get("/api/sessions/"+sid).get_json()["messages"])
    assert model.calls==["record_analysis_plan","list_data"]
    receipt=next(r for r in loop.context.analysis_state.tool_receipts if r["tool_name"]=="list_data")
    assert receipt["step_id"]=="empty_catalog" and receipt["plan_id"] and len(receipt["required_delivery_ids"])==2
    assert receipt["data_identities"]=={} and receipt["catalog_binding"]["session_id"]==sid
    exported=export_assistant_reply(sid,answer["body"],"md")
    assert answer["body"] in (cfg.sessions_resolved/sid/"reports"/exported["artifact_path"].split("/")[-1]).read_text(encoding="utf-8")


@pytest.mark.parametrize("phase",["synthesis","error","planning"])
def test_catalog_metadata_denied_phases_cannot_authorize_empty_facts(engine,phase):
    from data_agent.agent.execution_scope import WorkspaceScopeSnapshot
    from data_agent.agent.publication.native import resolve_native_claims
    loop,_=engine
    with loop.context.bind_workspace_scope(WorkspaceScopeSnapshot(phase=phase,session_id=loop.session_id)):
        rid,_=execute(loop,"list_data",{})
    with pytest.raises(ValueError):resolve_native_claims(session_id=loop.session_id,receipt_id=rid)


def test_catalog_metadata_replays_exact_scope_without_leaking_hidden_parent(engine):
    import pandas as pd
    from data_agent.agent.execution_scope import WorkspaceScopeSnapshot
    from data_agent.agent.publication.native import resolve_native_claims
    loop,_=engine
    _list_catalog_fixture(loop)
    with loop.context.bind_workspace_scope(WorkspaceScopeSnapshot(phase="execution",session_id=loop.session_id,allowed_datasets=frozenset({"derived_11"}))):
        rid,result=execute(loop,"list_data",{})
    loop.context.workspace.add("unrelated_hidden",pd.DataFrame({"secret":[1]}))
    claims=resolve_native_claims(session_id=loop.session_id,receipt_id=rid)
    assert claims[0]["value"]==result["datasets"]
    assert set(result["datasets"])=={"derived_11"} and result["datasets"]["derived_11"]["derived_from"] is None
    assert "unrelated_hidden" not in encode(claims)


def test_catalog_metadata_does_not_allow_empty_identity_for_other_tool(engine):
    from data_agent.agent.publication.facts import load_receipt_result
    loop,_=engine
    rid,_=execute(loop,"describe_dataset",{"name":"orders"})
    receipt=loop.context.analysis_state.tool_receipts[-1]
    receipt["data_identities"]={};receipt["dataset_refs"]=[]
    receipt["catalog_binding"]=loop.context.workspace._publication_catalog_result()["binding"]
    with pytest.raises(ValueError):load_receipt_result(session_id=loop.session_id,receipt_id=rid)


@pytest.mark.parametrize("scope",[{"datasets":[]},{"parameters":{"catalog_output":"datasets"},"datasets":[]},{"parameters":{"catalog_scope":"current_visible"},"datasets":[]}])
def test_catalog_metadata_required_table_scope_needs_explicit_selector_and_output(scope):
    from data_agent.agent.publication.deliveries import validate_business_scope
    with pytest.raises(ValueError):validate_business_scope(scope,required=True,metadata=True)


@pytest.mark.parametrize("case", ["persist", "meta", "state"])
def test_normal_app_sql_persistence_failure_has_no_ghost_dataset_after_restart(engine, tmp_path, monkeypatch, case):
    import sqlite3
    from data_agent.web.app import create_app
    from data_agent.web.agent_manager import AgentManager
    from data_agent.session.history import load_session
    from data_agent.agent.publication.runtime import read_answer
    from data_agent.tools.report import export_assistant_reply
    from data_agent.llm.client import Response, ToolCall
    _, cfg = engine
    database = tmp_path / "refusal.sqlite"
    with sqlite3.connect(database) as conn:
        conn.execute("CREATE TABLE incoming(amount INTEGER)")
        conn.executemany("INSERT INTO incoming VALUES (?)", [(7,), (13,), (22,)])
    connection = "sqlite:///" + str(database)
    query = "SELECT * FROM incoming"
    content = "从指定数据库查询载入experiment，报告本次查询记录数。"
    class DeniedModel(_NativeProductModel):
        def chat(self, messages, system="", **kw):
            from scripts.testing.run_p7_browser_rehearsal import ScriptedClient
            if system == "你是一个JSON输出器，只返回JSON，不要其他内容。":
                return Response(text=encode({"intent_type": "directed_analysis", "requested_outputs": [{"kind": "result", "exact_span": content, "business_scope": {"datasets": ["experiment"], "quantity": "rows"}}]}))
            if self.step == 0:
                self.step += 1; self.calls.append("load_sql")
                return Response(tool_calls=[ToolCall("sql_refusal", "load_sql", {"name": "experiment", "connection_string": connection, "query": query})])
            view = ScriptedClient._current_host_snapshot(messages, system)
            return Response(text=encode({"schema": "answer_bindings.v2", "snapshot_digest": view["snapshot_digest"], "sections": []}))
    model = DeniedModel(content, "experiment", "", {}, [])
    manager = AgentManager(client_factory=lambda **kw: model, auxiliary_client_factory=lambda: model)
    app = create_app(); app.config["agent_manager"] = manager; sid = "s3_sql_refusal_" + case
    from data_agent.session.workspace import Workspace
    from data_agent.agent.analysis_state import AnalysisSessionState
    owner,method=(AnalysisSessionState,"save") if case=="state" else (Workspace,"save_meta") if case=="meta" else (Workspace,"persist_dataset")
    original=getattr(owner,method);attempted=[]
    def fail(*a,**kw):
        import inspect
        value=original(*a,**kw)
        if any(f.function=="commit" and f.filename.endswith("data_io.py") for f in inspect.stack()):
            attempted.append(True)
            raise OSError("injected SQL ingestion durable write failure")
        return value
    monkeypatch.setattr(owner,method,fail)
    with app.test_client() as http:
        response = http.post("/api/chat", json={"session_id": sid, "message": content})
        events = [(b.split("\n", 1)[0][7:], json.loads(b.split("\ndata: ", 1)[1])) for b in response.get_data(as_text=True).split("\n\n") if b.startswith("event:")]
        ends = [d for e, d in events if e == "turn_end"]
        assert len(ends) == 1 and ends[0]["status"] == "completed", events
        saved = load_session(sid)["messages"][-1]; answer = read_answer(sid, saved["answer_ref"])
        assert answer["body"] == saved["content"] == ends[0]["reply_content"]
        assert answer["requirement_outcomes"][0]["status"] in {"failed","evidence_insufficient"}, answer
        assert attempted
        assert "secret" not in answer["body"] and not answer["source_snapshot"]["claims"]
        assert model.calls == ["load_sql"]
        refreshed = http.get("/api/sessions/" + sid).get_json()
        assert any(m.get("answer_ref") == saved["answer_ref"] for m in refreshed["messages"])
    with sqlite3.connect(database) as conn:
        assert conn.execute("SELECT amount FROM incoming ORDER BY amount").fetchall() == [(7,), (13,), (22,)]
    exported = export_assistant_reply(sid, answer["body"], "md")
    assert answer["body"] in (cfg.sessions_resolved / sid / "reports" / exported["artifact_path"].split("/")[-1]).read_text(encoding="utf-8")

    from data_agent.agent.context import use_agent_context
    for current in [manager.get(sid), AgentManager(client_factory=lambda **kw:model,auxiliary_client_factory=lambda:model).get_or_create(sid)]:
        with use_agent_context(current.context):
            assert not current.context.workspace.exists("experiment")
            assert not current.context.workspace.exists("experiment__raw")
        assert not current.context.analysis_state.dataset_contracts
    data_dir=cfg.sessions_resolved/sid/"data"
    assert not data_dir.exists() or not list(data_dir.iterdir())


@pytest.mark.parametrize("tool", ["export_data", "export_output"])
@pytest.mark.parametrize("case", ["nanoseconds", "aware", "microseconds", "time", "duration", "native_datetime", "native_date", "native_time", "missing", "bytes", "complex", "period"])
def test_excel_temporal_roundtrip_and_undefined_type_refusal(engine, tmp_path, tool, case):
    from datetime import datetime,date,time
    from openpyxl import load_workbook
    from data_agent.tools.registry import registry
    import pandas as pd
    loop,cfg=engine
    values={"nanoseconds":pd.Timestamp("2026-10-11T12:34:56.123456789"),"aware":pd.Timestamp("2026-10-11T12:34:56.123456789+08:00"),"microseconds":datetime(2026,10,11,12,34,56,123456),"time":time(12,34,56,123456),"duration":pd.Timedelta("1 days 00:00:00.123456789"),"native_datetime":datetime(2026,10,11,12,34,56,123000),"native_date":date(2026,10,11),"native_time":time(12,34,56,123000),"missing":None,"bytes":b"abc","complex":1+2j,"period":pd.Period("2026-10",freq="M")}
    value=values[case];frame=pd.DataFrame({"value":pd.Series([value],dtype=object)})
    loop.context.workspace.add("temporal",frame)
    destination=cfg.workspace_resolved/"data/temporal.xlsx";destination.parent.mkdir(parents=True,exist_ok=True);destination.write_bytes(b"old destination")
    args={"name":"temporal","path":str(destination),"fmt":"excel"}
    if tool=="export_output":args["output_type"]="data"
    result=registry.execute(tool,args)
    if case in {"bytes","complex","period"}:
        assert result.data["error_type"]=="tool_execution_error" and "Excel导出拒绝" in result.data["error"]
        assert destination.read_bytes()==b"old destination"
        assert not list((cfg.sessions_resolved/loop.session_id/"output").glob("export_*"))
        return
    info=json.loads(result.summary);book=load_workbook(destination,data_only=False);cell=book.active.cell(2,1)
    text=case in {"nanoseconds","aware","microseconds","time","duration"}
    if text:assert cell.value==value.isoformat() and cell.data_type=="s"
    elif case=="missing":assert cell.value is None
    elif case=="native_date":assert cell.value.date()==value and cell.data_type=="d"
    else:assert cell.value==value and cell.data_type=="d"
    audit=info["format_audit"]
    assert audit["temporal_as_text_total"]==int(text)
    assert audit["columns"][0]["temporal_as_text_count"]==int(text)
    if text:assert "时间或时长转文本" in info["message"]
    book.close();pd.testing.assert_frame_equal(loop.context.workspace.get("temporal"),frame)


def _export_failure(monkeypatch, destination, output, fmt, phase):
    """Inject one storage/encoding failure; retain actual partial bytes."""
    from pathlib import Path
    from data_agent.tools import data_io
    from data_agent.session import history
    import pandas as pd
    fired = []
    previous_destination = destination.read_bytes() if destination.exists() else None
    if phase == "serialize":
        if fmt == "csv":
            original = pd.DataFrame.to_csv
            def fail(frame, path_or_buf=None, *args, **kwargs):
                if not fired:
                    fired.append(phase)
                    if isinstance(path_or_buf, (str, Path)):
                        Path(path_or_buf).write_bytes(b"partial CSV")
                    elif path_or_buf is not None:
                        path_or_buf.write(b"partial CSV")
                    raise OSError("injected export serialize failure")
                return original(frame, path_or_buf, *args, **kwargs)
            monkeypatch.setattr(pd.DataFrame, "to_csv", fail)
        else:
            method = "_serialize_json_export" if fmt == "json" else "_serialize_excel_export"
            original = getattr(data_io, method)
            def fail(*args, **kwargs):
                value = original(*args, **kwargs)
                if not fired:
                    fired.append(phase)
                    raise OSError("injected export serialize failure")
                return value
            monkeypatch.setattr(data_io, method, fail)
    elif phase in {"temporary", "snapshot"}:
        original = Path.write_bytes
        def fail(path, data):
            committed_destination = destination.exists() and destination.read_bytes() != previous_destination
            selected = (path.parent == destination.parent and path.name.startswith(".atomic-")) if phase == "temporary" else (committed_destination and path.parent == output and (path.name.startswith("export_") or path.name.startswith(".atomic-")))
            if selected and not fired:
                fired.append(phase)
                original(path, b"partial export bytes")
                raise OSError("injected export " + phase + " failure")
            return original(path, data)
        monkeypatch.setattr(Path, "write_bytes", fail)
    elif phase in {"replace", "after_replace"}:
        original = Path.replace
        def fail(path, target):
            if Path(target).resolve() == destination.resolve() and not fired:
                fired.append(phase)
                if phase == "after_replace": original(path, target)
                raise OSError("injected export " + phase + " failure")
            return original(path, target)
        monkeypatch.setattr(Path, "replace", fail)
    else:
        assert phase == "manifest"
        original = history.register_artifact
        def fail(*args, **kwargs):
            value = original(*args, **kwargs)
            if not fired:
                fired.append(phase)
                raise OSError("injected export manifest failure")
            return value
        monkeypatch.setattr(history, "register_artifact", fail)
    return fired


@pytest.mark.parametrize("tool", ["export_data", "export_output"])
@pytest.mark.parametrize("fmt", ["csv", "json", "excel"])
@pytest.mark.parametrize("existing", [False, True])
@pytest.mark.parametrize("phase", ["serialize", "temporary", "replace", "after_replace", "snapshot", "manifest"])
def test_export_failure_preserves_previous_bytes_artifacts_and_source(engine, monkeypatch, tool, fmt, existing, phase):
    from pathlib import Path
    from data_agent.tools.registry import registry
    import pandas as pd
    loop, cfg = engine
    destination = cfg.workspace_resolved / "data" / ("existing." + {"excel": "xlsx"}.get(fmt, fmt))
    destination.parent.mkdir(parents=True, exist_ok=True)
    args = {"name": "orders", "path": str(destination), "fmt": fmt}
    if tool == "export_output": args["output_type"] = "data"
    own = cfg.sessions_resolved / loop.session_id
    if existing:
        prior = registry.execute(tool, args)
        assert not (prior.data or {}).get("error"), prior.summary
    before = {p.relative_to(own).as_posix(): p.read_bytes() for p in own.rglob("*") if p.is_file()}
    previous = destination.read_bytes() if existing else None
    changed = loop.context.workspace.get("orders").copy(deep=True)
    changed["amount"] += 1
    loop.context.workspace.add("orders", changed)
    frame = loop.context.workspace.get("orders").copy(deep=True)
    identity = loop.context.workspace.get_data_identity("orders")
    with monkeypatch.context() as fault:
        fired = _export_failure(fault, destination, own / "output", fmt, phase)
        result = registry.execute(tool, args)
    assert fired == [phase]
    assert (result.data or {}).get("error"), result.summary
    assert destination.read_bytes() == previous if existing else not destination.exists()
    assert {p.relative_to(own).as_posix(): p.read_bytes() for p in own.rglob("*") if p.is_file()} == before
    assert not list(destination.parent.glob(".atomic-*.tmp"))
    assert loop.context.workspace.get_data_identity("orders") == identity
    pd.testing.assert_frame_equal(loop.context.workspace.get("orders"), frame)


@pytest.mark.parametrize('tool,fmt,phase', [('export_data', 'csv', 'temporary'), ('export_data', 'json', 'snapshot'), ('export_data', 'excel', 'manifest'), ('export_output', 'csv', 'manifest'), ('export_output', 'json', 'temporary'), ('export_output', 'excel', 'snapshot')])
def test_normal_app_failed_export_keeps_old_download_after_version_change_refresh_and_restart(engine, monkeypatch, tool, fmt, phase):
    import pandas as pd
    from data_agent.web.app import create_app
    from data_agent.web.agent_manager import AgentManager
    from data_agent.session.history import load_session
    from data_agent.agent.publication.runtime import read_answer
    from data_agent.tools.report import export_assistant_reply
    from data_agent.agent.context import use_agent_context
    from data_agent.llm.client import Response, ToolCall
    from data_agent.tools.data_io import _record_trust_workflow
    _, cfg = engine
    content = "只将当前experiment完整导出为指定格式文件，提供本次下载。"
    args = {"name": "experiment", "path": "output/current." + {"excel": "xlsx"}.get(fmt, fmt), "fmt": fmt}
    if tool == "export_output": args["output_type"] = "data"
    class ExportModel(_NativeProductModel):
        def chat(self, messages, system="", **kw):
            from scripts.testing.run_p7_browser_rehearsal import ScriptedClient
            if system == "你是一个JSON输出器，只返回JSON，不要其他内容。":
                return Response(text=encode({"intent_type": "data_operation", "requested_outputs": [{"kind": "result", "exact_span": self.content, "business_scope": {"datasets": ["experiment"]}}]}))
            if self.step == 0:
                self.step = 1; self.calls.append("quick_profile")
                return Response(tool_calls=[ToolCall("export_profile", "quick_profile", {"name": "experiment"})])
            if self.step < 3: return super().chat(messages, system=system, **kw)
            view = ScriptedClient._current_host_snapshot(messages, system)
            return Response(text=encode({"schema": "answer_bindings.v2", "snapshot_digest": view["snapshot_digest"], "sections": [{"requirement_ids": [r["id"]], "parts": [{"kind": "artifact", "ref": a} for a in r.get("artifact_ids", [])]} for r in view["requirements"] if not r["withdrawn"]]}))
    model = ExportModel(content, "experiment", tool, args, [{"kind": "artifact", "scope": {"datasets": ["experiment"]}}])
    manager = AgentManager(client_factory=lambda **kw: model, auxiliary_client_factory=lambda: model)
    sid = "s3_atomic_" + tool + "_" + fmt + "_" + phase
    loop = manager.get_or_create(sid)
    def install(values):
        frame = pd.DataFrame({"amount": values, "text": ["中文", "original"]})
        with use_agent_context(loop.context):
            loop.context.workspace.add("experiment", frame)
            _record_trust_workflow(session_id=sid, state=loop.context.analysis_state, dataset="experiment", df=frame, applied=[], needs_confirm=[], quality={}, interpretation_data={}, detail_path="")
    install([10, 20])
    app = create_app(); app.config["agent_manager"] = manager
    def turn(http):
        response = http.post("/api/chat", json={"session_id": sid, "message": model.content})
        events = [(b.split("\n", 1)[0][7:], json.loads(b.split("\ndata: ", 1)[1])) for b in response.get_data(as_text=True).split("\n\n") if b.startswith("event:")]
        ends = [d for e, d in events if e == "turn_end"]
        assert len(ends) == 1 and ends[0]["status"] == "completed", events
        saved = load_session(sid)["messages"][-1]
        answer = read_answer(sid, saved["answer_ref"])
        assert answer["body"] == saved["content"] == ends[0]["reply_content"]
        return saved, answer
    own = cfg.sessions_resolved / sid; working = own / "output" / ("current." + {"excel": "xlsx"}.get(fmt, fmt))
    with app.test_client() as http:
        first_saved, first = turn(http)
        assert first["requirement_outcomes"][-1]["status"] == "answered", first
        artifact = first["source_snapshot"]["artifacts"][0]
        download_url = "/api/files/sessions/" + sid + "/" + artifact["path"]
        old_download = http.get(download_url).data; old_working = working.read_bytes()
        manifest = (own / "artifacts.json").read_bytes()
        published = {p.name: p.read_bytes() for p in (own / "output").glob("export_*")}
        prior_receipts = {r["id"] for r in loop.context.analysis_state.tool_receipts if r["tool_name"] == tool}
        install([30, 40]); current_identity = loop.context.workspace.get_data_identity("experiment")
        model.content = "将当前experiment新版重新完整导出为指定格式文件，提供本次下载。"
        model.step = 0; model.calls.clear()
        with monkeypatch.context() as fault:
            fired = _export_failure(fault, working, own / "output", fmt, phase)
            saved, answer = turn(http)
        assert fired == [phase]
        assert answer["requirement_outcomes"][-1]["status"] != "answered", answer
        assert not answer["artifact_refs"] and not answer["source_snapshot"]["artifacts"]
        assert working.read_bytes() == old_working
        current_manifest = json.loads((own / "artifacts.json").read_bytes())
        previous_manifest = json.loads(manifest)
        assert all(entry in current_manifest for entry in previous_manifest)
        assert [entry for entry in current_manifest if entry["type"] == "file"] == [entry for entry in previous_manifest if entry["type"] == "file"]
        additions = [entry for entry in current_manifest if entry not in previous_manifest]
        assert sorted(entry["type"] for entry in additions) == ["accepted_answer", "analysis_plan"]
        assert {p.name: p.read_bytes() for p in (own / "output").glob("export_*")} == published
        assert not list((own / "output").glob(".atomic-*.tmp"))
        assert http.get(download_url).data == old_download
        assert read_answer(sid, first_saved["answer_ref"]) == first
        assert model.calls == ["quick_profile", "record_analysis_plan", tool]
        assert {r["id"] for r in loop.context.analysis_state.tool_receipts if r["tool_name"] == tool} == prior_receipts
        assert loop.context.workspace.get_data_identity("experiment") == current_identity
        refreshed = http.get("/api/sessions/" + sid).get_json()
        assert {m.get("answer_ref", {}).get("path") for m in refreshed["messages"] if m.get("answer_ref")} >= {first_saved["answer_ref"]["path"], saved["answer_ref"]["path"]}
        exported = export_assistant_reply(sid, answer["body"], "md")
        assert answer["body"] in (own / "reports" / exported["artifact_path"].split("/")[-1]).read_text(encoding="utf-8")
    restarted = AgentManager(client_factory=lambda **kw: model, auxiliary_client_factory=lambda: model)
    app.config["agent_manager"] = restarted
    with app.test_client() as http:
        assert http.get(download_url).data == old_download
        assert any(m.get("answer_ref") == saved["answer_ref"] for m in http.get("/api/sessions/" + sid).get_json()["messages"])
    assert read_answer(sid, first_saved["answer_ref"]) == first


@pytest.mark.parametrize("n", [-1, 0, 51, 100, True, 1.5, "10", None])
def test_preview_rejects_unsupported_requested_row_count_before_reading(engine, monkeypatch, n):
    from data_agent.tools import data_understand
    from data_agent.tools.registry import registry
    called = []
    monkeypatch.setattr(data_understand, "get_df", lambda name: called.append(name))
    result = registry.execute("preview_data", {"name": "orders", "n": n})
    assert (result.data or {}).get("error"), result.summary
    assert called == []


@pytest.mark.parametrize("case,n,total,expected", [("default", None, 60, 10), ("one", 1, 60, 1), ("ten", 10, 60, 10), ("limit", 50, 60, 50), ("short", 50, 3, 3), ("empty", 10, 0, 0)])
def test_preview_original_values_order_and_explicit_requested_returned_source_scope(engine, case, n, total, expected):
    import pandas as pd
    from data_agent.tools.registry import registry
    loop, _ = engine
    frame = pd.DataFrame({"row_id": range(total), "amount": [0.423456789012345] * total})
    loop.context.workspace.add("preview_source", frame)
    identity = loop.context.workspace.get_data_identity("preview_source")
    args = {"name": "preview_source"}
    if n is not None: args["n"] = n
    result = registry.execute("preview_data", args)
    data = result.data
    assert not (data or {}).get("error"), result.summary
    assert data["data"] == frame.head(expected).to_dict("records")
    assert data["rows"] == expected and data["shape"] == {"rows": total, "columns": 2}
    assert data["observation"] == {"requested_rows": 10 if n is None else n, "returned_rows": expected, "source_rows": total, "preview_limit": 50, "row_order": "current_dataset_order", "scope": "first_rows"}
    assert "预览" in result.summary and "请求" in result.summary and "实际" in result.summary and "源表" in result.summary
    pd.testing.assert_frame_equal(loop.context.workspace.get("preview_source"), frame)
    assert loop.context.workspace.get_data_identity("preview_source") == identity


@pytest.mark.parametrize("fault", ["value", "order", "rows", "columns", "shape", "requested", "returned", "total"])
def test_preview_closed_adapter_rejects_changed_records_or_scope(engine, monkeypatch, fault):
    from copy import deepcopy
    from data_agent.agent.publication import native
    loop, _ = engine
    rid, _ = execute(loop, "preview_data", {"name": "orders", "n": 2})
    original = native.load_receipt_result
    def changed(**kwargs):
        context, receipt, result, binding = original(**kwargs); result = deepcopy(result)
        if fault == "value": result["data"][0]["amount"] = 999
        elif fault == "order": result["data"].reverse()
        elif fault == "rows": result["rows"] = 99
        elif fault == "columns": result["columns"] = ["other"]
        elif fault == "shape": result["shape"]["rows"] = 999
        else: result.setdefault("observation", {})[{"requested": "requested_rows", "returned": "returned_rows", "total": "source_rows"}[fault]] = 999
        return context, receipt, result, binding
    monkeypatch.setattr(native, "load_receipt_result", changed)
    with pytest.raises(ValueError): native.resolve_native_claims(session_id=loop.session_id, receipt_id=rid)


@pytest.mark.parametrize("phase", ["planning", "synthesis", "error"])
def test_preview_cannot_read_record_values_in_forbidden_scope(engine, phase):
    from data_agent.agent.execution_scope import WorkspaceScopeSnapshot
    from data_agent.tools.registry import registry
    loop, _ = engine
    with loop.context.bind_workspace_scope(WorkspaceScopeSnapshot(phase=phase)):
        result = registry.execute("preview_data", {"name": "orders", "n": 2})
    assert (result.data or {}).get("error"), result.summary
    assert not (result.data or {}).get("data")


@pytest.mark.parametrize("requested", [1, 2])
def test_preview_delivery_matches_exact_requested_row_count(engine, requested):
    from data_agent.agent.publication.runtime import snapshot_for_turn
    loop, _ = engine
    request(loop, "交付当前前两行预览", [{"kind": "table", "scope": {"datasets": ["orders"], "parameters": {"preview_n": requested, "preview_output": "data"}}}])
    execute(loop, "preview_data", {"name": "orders", "n": 2})
    req = snapshot_for_turn(loop).as_dict()["requirements"][0]
    assert bool(req["claim_ids"]) == (requested == 2)


@pytest.mark.parametrize('case', ['default', 'limit', 'empty', 'rejected'])
def test_normal_app_preview_actual_upload_formal_plan_scope_sse_save_refresh_and_export(engine, case):
    import io
    import pandas as pd
    from data_agent.web.app import create_app
    from data_agent.web.agent_manager import AgentManager
    from data_agent.session.history import load_session
    from data_agent.agent.publication.runtime import read_answer
    from data_agent.tools.report import export_assistant_reply
    _, cfg = engine
    total = 0 if case == "empty" else 3 if case == "short" else 60
    n = 1 if case == "one" else 50 if case in {"limit", "short"} else 100 if case == "rejected" else 10
    expected = min(n, total)
    frame = pd.DataFrame({"row_id": range(total), "amount": range(total)})
    content = "预览上传数据，将其命名为experiment；按源表当前顺序提供前" + str(n) + "行，并说明请求行数、实际返回数和源表总行数，不能把预览说成全表交付。"
    params = {} if case == "rejected" else {"preview_n": n}
    specs = [{"kind": "table", "scope": {"datasets": ["experiment"], "parameters": dict(params, preview_output=output)}} for output in ["data", "observation"]]
    args = {"name": "experiment"}
    if case != "default": args["n"] = n
    model = _NativeProductModel(content, "experiment", "preview_data", args, specs)
    manager = AgentManager(client_factory=lambda **kw: model, auxiliary_client_factory=lambda: model)
    app = create_app(); app.config["agent_manager"] = manager
    sid = "s3_preview_" + case; raw = frame.to_csv(index=False).encode("utf-8")
    with app.test_client() as http:
        ticket = http.post("/api/upload", data={"file": (io.BytesIO(raw), "preview.csv")}, content_type="multipart/form-data").get_json()
        assert ticket["sha256"] == "sha256:" + hashlib.sha256(raw).hexdigest(); model.ticket = ticket
        response = http.post("/api/chat", json={"session_id": sid, "message": content, "uploads": [ticket]})
        events = [(b.split("\n", 1)[0][7:], json.loads(b.split("\ndata: ", 1)[1])) for b in response.get_data(as_text=True).split("\n\n") if b.startswith("event:")]
        ends = [d for e, d in events if e == "turn_end"]
        assert len(ends) == 1 and ends[0]["status"] == "completed", events
        refreshed = http.get("/api/sessions/" + sid).get_json()
    saved = load_session(sid)["messages"][-1]; answer = read_answer(sid, saved["answer_ref"])
    assert answer["body"] == saved["content"] == ends[0]["reply_content"]
    assert answer["requirement_outcomes"][0]["status"] == ("not_computed" if case == "rejected" else "answered"), answer
    if case == "rejected":
        assert not any(c["target"]["tool"] == "preview_data" for c in answer["source_snapshot"]["claims"])
        assert not any(r["tool_name"] == "preview_data" for r in manager.get(sid).context.analysis_state.tool_receipts)
    else:
        values = {c["basis_refs"][0]["result_path"][0]: c["value"] for c in answer["source_snapshot"]["claims"] if c["target"]["tool"] == "preview_data"}
        assert values["data"] == frame.head(expected).to_dict("records")
        assert values["observation"]["requested_rows"] == n and values["observation"]["returned_rows"] == expected and values["observation"]["source_rows"] == total
        assert "预览" in answer["body"] and "全表" in answer["body"]
        receipt = next(r for r in manager.get(sid).context.analysis_state.tool_receipts if r["tool_name"] == "preview_data")
        assert receipt["plan_id"] and receipt["step_id"] == "native_product" and len(receipt["required_delivery_ids"]) == 2
    assert model.calls == ["describe_dataset", "record_analysis_plan", "preview_data"]
    assert any(m.get("answer_ref") == saved["answer_ref"] for m in refreshed["messages"])
    exported = export_assistant_reply(sid, answer["body"], "md")
    assert answer["body"] in (cfg.sessions_resolved / sid / "reports" / exported["artifact_path"].split("/")[-1]).read_text(encoding="utf-8")

